"""Tests for original-record Word export."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from openpyxl import Workbook

from src.generators.original_record import (
    OriginalRecordData,
    build_original_record_document,
    default_output_path,
    export_leg_original_record,
    export_node_original_record,
    format_id_range,
    format_share_path_for_record,
    format_test_period,
    generate_original_record,
    original_record_data_from_node,
    resolve_combined_original_record_folder,
    resolve_original_record_folder,
    resolve_original_record_template,
    stack_standard_blocks,
)
from src.io.test_photos import TEST_GROUP_DIR
from src.models.project_state import (
    DataTableRef as DataRef,
    ProjectState,
    TestEquipment as EqModel,
    TestNode,
    TestSample,
    TestStandard as StdModel,
)

TEMPLATE = Path("templates/original_data_sheet/original_record_placeholders.docx")


def test_resolve_template_falls_back_local():
    path = resolve_original_record_template()
    assert path.is_file()
    assert path.name == "original_record_placeholders.docx"


def test_format_helpers():
    assert format_id_range(["A01", "A02", "A03"]) == "A01~A03"
    assert format_id_range(["A01", "B02"]) == "A01、B02"
    assert format_test_period("2026-09-08", "2026-09-09") == "2026.09.08~2026.09.09"
    assert format_share_path_for_record(
        "/Volumes/材料实验室b/车载电子/2026年/foo",
        smb_hint="smb://10.10.31.8/材料实验室b/车载电子",
    ) == r"\\10.10.31.8\材料实验室b\车载电子\2026年\foo"


def test_stack_standard_blocks():
    stds = [
        StdModel(test_name="盐雾", standard_desc="条件A", evaluation_req="要求A"),
        StdModel(test_name="湿热", standard_desc="条件B", evaluation_req="要求B"),
    ]
    text = stack_standard_blocks(stds, "standard_desc")
    assert "盐雾" in text and "条件A" in text
    assert "湿热" in text and "条件B" in text


def test_resolve_original_record_folder_prefers_remote(tmp_path: Path):
    local = tmp_path / "local"
    remote = tmp_path / "remote"
    local.mkdir()
    remote.mkdir()
    folder = resolve_original_record_folder(
        local, remote, leg_name="Leg 1", test_item="沙尘试验"
    )
    assert folder == remote / TEST_GROUP_DIR / "Leg 1-沙尘试验"


def test_resolve_original_record_folder_falls_back_local(tmp_path: Path):
    local = tmp_path / "local"
    local.mkdir()
    folder = resolve_original_record_folder(
        local, tmp_path / "missing", leg_name="Leg 1", test_item="沙尘试验"
    )
    assert folder == local / TEST_GROUP_DIR / "Leg 1-沙尘试验"


def test_resolve_original_record_folder_requires_roots_and_names(tmp_path: Path):
    root = tmp_path / "proj"
    root.mkdir()
    assert resolve_original_record_folder(None, None, leg_name="Leg 1", test_item="沙尘") is None
    assert resolve_original_record_folder(root, None, leg_name="", test_item="沙尘") is None
    assert (
        resolve_original_record_folder(root, None, leg_name="Leg 1", test_item="请选择试验...")
        is None
    )


def test_original_record_data_from_node_uses_persisted_fields():
    node = TestNode(
        test_name="沙尘试验",
        test_method="VW 1",
        env_condition="23C",
        start_date="2026-09-08",
        end_date="2026-09-09",
    )
    node.apply_standards([StdModel(standard_id="VW", chapter="8", test_name="沙尘试验")])
    node.equipments = [EqModel(name="沙尘箱", code="T1")]
    node.samples = [TestSample(sample_id="A01"), TestSample(sample_id="  ")]
    node.data_tables = [DataRef(title="工况", relative_path="a.xlsx")]
    state = ProjectState(
        project_id="P1",
        sample_name="控制器",
        tester_name="张三",
        source_path="/Volumes/share/proj",
        project_path="/tmp/local",
        application_fields={"申请单号": "A226"},
    )
    data = original_record_data_from_node(node, leg_name="Leg 1", state=state)
    assert data.application_no == "A226"
    assert data.test_item == "沙尘试验"
    assert data.test_method == "VW 1"
    assert data.sample_name == "控制器"
    assert data.sample_ids == ["A01"]
    assert data.env_condition == "23C"
    assert data.start_date == "2026-09-08"
    assert data.end_date == "2026-09-09"
    assert data.share_path == "/Volumes/share/proj"
    assert data.equipments == node.equipments
    assert [s.standard_id for s in data.standards] == ["VW"]
    assert data.tester_name == "张三"
    assert data.data_tables == node.data_tables
    assert data.project_path == "/tmp/local"
    assert data.remote_root == "/Volumes/share/proj"
    assert data.leg_name == "Leg 1"


def test_export_node_original_record_requires_project_dir():
    data = OriginalRecordData(test_item="沙尘", leg_name="Leg 1")
    try:
        export_node_original_record(data, template_path=TEMPLATE)
    except FileNotFoundError as exc:
        assert "项目目录" in str(exc)
        return
    assert False, "expected FileNotFoundError"


def test_default_output_path_saves_beside_albums(tmp_path: Path):
    folder = tmp_path / TEST_GROUP_DIR / "Leg 1-沙尘试验"
    (folder / "试验前").mkdir(parents=True)
    path = default_output_path(folder, application_no="A1", test_item="沙尘试验")
    assert path.parent == folder
    assert path.name == "A1_沙尘试验_原始记录.docx"
    assert not (tmp_path / "原始记录").exists()
    assert not (folder / "原始记录").exists()
    first = default_output_path(folder, application_no="A1", test_item="沙尘试验")
    first.write_bytes(b"x")
    second = default_output_path(folder, application_no="A1", test_item="沙尘试验")
    assert second.name == "A1_沙尘试验_原始记录-2.docx"


def test_generate_original_record_fills_and_expands(tmp_path: Path):
    assert TEMPLATE.is_file()
    data = OriginalRecordData(
        application_no="A22606909401",
        test_item="Dust test",
        test_method="VW82511-2010 / 8.3.6",
        sample_name="控制器",
        sample_ids=["A22606909401-A01", "A22606909401-A02"],
        env_condition="23℃ / 50%RH",
        start_date="2026-09-08",
        end_date="2026-09-08",
        share_path=r"\\10.10.31.8\材料实验室b\车载电子\proj",
        equipments=[
            EqModel(name="沙尘箱", model="M1", code="TTE-001", valid_date="2027-01-01"),
            EqModel(name="温箱", model="M2", code="TTE-002", valid_date="2027-02-01"),
        ],
        standards=[
            StdModel(
                test_name="沙尘试验",
                standard_desc="粉尘浓度 …",
                evaluation_req="外观无异常",
            ),
            StdModel(
                test_name="功能检查",
                standard_desc="通电检查",
                evaluation_req="功能正常",
            ),
        ],
        tester_name="张三",
    )
    out = tmp_path / "out.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    assert out.is_file()

    doc = Document(str(out))
    blob = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for tc in row._tr.tc_lst:
                blob += "\n" + _Cell(tc, table).text

    assert "A22606909401" in blob
    assert "Dust test" in blob
    assert "控制器" in blob
    assert "A22606909401-A01~A22606909401-A02" in blob
    assert "23℃ / 50%RH" in blob
    assert "张三" in blob
    assert "沙尘箱" in blob and "温箱" in blob
    assert "沙尘试验" in blob and "功能检查" in blob
    assert "{{申请单编号}}" not in blob
    assert "{{设备名称}}" not in blob
    assert "{{试验项目}}" not in blob
    assert "{{结果样品编号}}" not in blob
    assert "{{RESULT_TABLE}}" not in blob
    assert "{{DATA_TABLES}}" not in blob
    assert "{{" not in blob
    assert any(p.text.strip() == "/" for p in doc.paragraphs)

    resultish = [
        t
        for t in doc.tables
        if any(
            "试验前" in _Cell(tc, t).text
            for row in t.rows
            for tc in row._tr.tc_lst
        )
    ]
    assert len(resultish) == 2
    merge_cols = _template_result_merge_cols()
    for t in resultish:
        starts = sum(
            1
            for row in t.rows
            for tc in row._tr.tc_lst
            if _Cell(tc, t).text.strip() == "试验前"
            or _Cell(tc, t).text.strip().startswith("试验前")
        )
        assert starts == 2
        _assert_sample_block_vmerges(t, merge_cols)


def _vmerge_val(tc) -> str | None:
    tcPr = tc.tcPr
    if tcPr is None:
        return None
    vm = tcPr.find(qn("w:vMerge"))
    if vm is None:
        return None
    return vm.get(qn("w:val")) or "continue"


def _result_sample_starts(table) -> list[int]:
    starts: list[int] = []
    for i, row in enumerate(table.rows):
        texts = [_Cell(tc, table).text.strip() for tc in row._tr.tc_lst]
        if any(t == "试验前" or t.startswith("试验前") for t in texts):
            starts.append(i)
    return starts


def _template_result_merge_cols() -> list[int]:
    doc = Document(str(TEMPLATE))
    table = next(
        t
        for t in doc.tables
        if any(
            "试验前" in _Cell(tc, t).text
            for row in t.rows
            for tc in row._tr.tc_lst
        )
    )
    start = _result_sample_starts(table)[0]
    cols = [
        ci
        for ci, tc in enumerate(table.rows[start]._tr.tc_lst)
        if _vmerge_val(tc) is not None
    ]
    # 样品编号 / 功能等级 / 结论
    assert cols[:1] == [0] and len(cols) >= 3
    return cols


def _assert_sample_block_vmerges(table, merge_cols: list[int]) -> None:
    starts = _result_sample_starts(table)
    assert starts
    for start in starts:
        for offset, expected in ((0, "restart"), (1, "continue"), (2, "continue")):
            tcs = table.rows[start + offset]._tr.tc_lst
            for ci in merge_cols:
                assert _vmerge_val(tcs[ci]) == expected, (
                    f"sample@{start} row+{offset} col {ci}: "
                    f"expected vMerge={expected}, got {_vmerge_val(tcs[ci])}"
                )


def test_generate_original_record_embeds_standard_images_like_report(tmp_path: Path):
    """测试参数 cell: text then blank then library images, same width as report."""
    from docx.shared import Inches

    from src.generators.word_engine import WordGenerator

    _MIN_PNG = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
        b"\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    from PIL import Image
    import io

    im = Image.new("RGB", (900, 300), (20, 20, 20))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    png = buf.getvalue()
    expected_w = WordGenerator._condition_image_width_in(png)

    data = OriginalRecordData(
        application_no="A1",
        test_item="振动",
        standards=[
            StdModel(
                test_name="机械冲击",
                standard_desc="冲击条件正文。",
                images=[png],
            ),
            StdModel(
                test_name="防尘实验",
                standard_desc="防尘条件正文。",
                images=[_MIN_PNG, png],
            ),
        ],
    )
    out = tmp_path / "with_imgs.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))

    cell = None
    for table in doc.tables:
        for row in table.rows:
            for tc in row._tr.tc_lst:
                c = _Cell(tc, table)
                if any("{{测试参数}}" in (p.text or "") for p in c.paragraphs):
                    assert False, "placeholder should be gone"
                texts = [p.text for p in c.paragraphs]
                if "机械冲击" in texts and "防尘实验" in texts:
                    cell = c
                    break
            if cell is not None:
                break
        if cell is not None:
            break
    assert cell is not None

    paras = list(cell.paragraphs)
    mech = next(i for i, p in enumerate(paras) if p.text == "机械冲击")
    dust = next(i for i, p in enumerate(paras) if p.text == "防尘实验")
    assert mech < dust
    assert any(_para_has_image(p) for p in paras[mech:dust])
    assert paras[mech + 2].text == ""
    assert _para_has_image(paras[mech + 3])

    dust_images = [i for i, p in enumerate(paras) if i > dust and _para_has_image(p)]
    assert len(dust_images) == 2
    assert paras[dust_images[0] - 1].text == ""
    assert paras[dust_images[1] - 1].text == ""

    widths = []
    for p in paras:
        for ext in p._element.xpath(".//a:ext"):
            cx = ext.get("cx")
            if cx:
                widths.append(int(cx))
    assert Inches(expected_w) in widths
    assert "{{测试参数}}" not in "\n".join(p.text for p in paras)


def _para_has_image(paragraph) -> bool:
    return bool(paragraph._element.xpath(".//w:drawing"))


def _write_simple_xlsx(path: Path, *, with_limit: bool, limit_value: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "样品编号"
    ws["B1"] = "读数"
    if with_limit:
        ws["A2"] = "上限"
        ws["B2"] = limit_value
        ws["A3"] = "A01"
        ws["B3"] = "1.0"
    else:
        ws["A2"] = "A01"
        ws["B2"] = "1.0"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def test_generate_original_record_data_tables_titles_and_limit_flag(tmp_path: Path):
    project = tmp_path / "proj"
    x1 = project / "t1.xlsx"
    x2 = project / "t2.xlsx"
    _write_simple_xlsx(x1, with_limit=True, limit_value="2")
    _write_simple_xlsx(x2, with_limit=False)

    refs = [
        DataRef(
            title="5点法",
            relative_path=str(x1.relative_to(project)),
            include_limit_row_in_report=True,
        ),
        DataRef(
            title="工况",
            relative_path=str(x2.relative_to(project)),
            include_limit_row_in_report=False,
        ),
    ]
    data = OriginalRecordData(
        application_no="A1",
        test_item="Dust",
        sample_ids=["A01"],
        project_path=str(project),
        leg_name="Leg1",
        data_tables=refs,
        standards=[StdModel(test_name="沙尘")],
    )
    out = tmp_path / "with_tables.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert "5点法" in paras
    assert "工况" in paras
    assert "{{DATA_TABLES}}" not in "\n".join(paras)

    joined = "\n---\n".join(
        "\n".join(_Cell(tc, t).text for row in t.rows for tc in row._tr.tc_lst)
        for t in doc.tables
    )
    assert "上限" in joined
    assert "2" in joined
    assert "读数" in joined
    assert "1.0" in joined


def test_generate_original_record_omits_limit_when_unchecked(tmp_path: Path):
    project = tmp_path / "proj"
    x1 = project / "t1.xlsx"
    _write_simple_xlsx(x1, with_limit=True, limit_value="2")
    refs = [
        DataRef(
            title="5点法",
            relative_path=str(x1.relative_to(project)),
            include_limit_row_in_report=False,
        )
    ]
    data = OriginalRecordData(
        test_item="Dust",
        project_path=str(project),
        data_tables=refs,
        standards=[StdModel(test_name="沙尘")],
    )
    out = tmp_path / "no_limit.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))
    joined = "\n".join(
        _Cell(tc, t).text for t in doc.tables for row in t.rows for tc in row._tr.tc_lst
    )
    assert "上限" not in joined
    assert "1.0" in joined


def _write_doubled_five_point_xlsx(path: Path) -> None:
    """Copy 5点法 header blocks twice → 1 sample + 18 content cols."""
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    temps = ["-40°C", "25°C", "85°C"] * 2
    volts = ["9V", "14V", "16V"]
    ws["A1"] = "样品编号 / Sample No."
    ws.merge_cells("A1:A2")
    col = 2
    for temp in temps:
        start = get_column_letter(col)
        end = get_column_letter(col + 2)
        ws.cell(1, col, temp)
        ws.merge_cells(f"{start}1:{end}1")
        for i, v in enumerate(volts):
            ws.cell(2, col + i, v)
        col += 3
    for r, sid in enumerate(["A01", "A02", "A03"], start=3):
        ws.cell(r, 1, f"A22600280178-{sid}")
        for c in range(2, col):
            ws.cell(r, c, str(r - 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def test_generate_original_record_data_table_equal_widths(tmp_path: Path):
    from src.generators.word_engine import _OR_SAMPLE_COL_DXA, _OR_TABLE_WIDTH_DXA

    project = tmp_path / "proj"
    x1 = project / "t1.xlsx"
    _write_simple_xlsx(x1, with_limit=False)
    refs = [
        DataRef(title="工况", relative_path=str(x1.relative_to(project))),
    ]
    data = OriginalRecordData(
        test_item="Dust",
        sample_ids=["A01"],
        project_path=str(project),
        data_tables=refs,
        standards=[StdModel(test_name="沙尘")],
    )
    out = tmp_path / "eq.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))

    # Find the generated data table (2 cols: 样品编号 + 读数)
    data_table = None
    for t in doc.tables:
        texts = [_Cell(tc, t).text for row in t.rows for tc in row._tr.tc_lst]
        if any("读数" in (x or "") for x in texts):
            data_table = t
            break
    assert data_table is not None
    grid = data_table._tbl.find(qn("w:tblGrid"))
    widths = [int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))]
    assert widths[0] == _OR_SAMPLE_COL_DXA
    assert sum(widths) == _OR_TABLE_WIDTH_DXA
    assert len(widths) == 2
    assert widths[1] == _OR_TABLE_WIDTH_DXA - _OR_SAMPLE_COL_DXA
    tblPr = data_table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    layout = tblPr.find(qn("w:tblLayout"))
    ind = tblPr.find(qn("w:tblInd"))
    assert tblW is not None and tblW.get(qn("w:w")) == str(_OR_TABLE_WIDTH_DXA)
    assert tblW.get(qn("w:type")) == "dxa"
    assert layout is not None and layout.get(qn("w:type")) == "fixed"
    assert ind is not None and ind.get(qn("w:w")) == "-431"


def test_generate_original_record_splits_wide_five_point_table(tmp_path: Path):
    from src.generators.word_engine import _OR_SAMPLE_COL_DXA, _OR_TABLE_WIDTH_DXA

    project = tmp_path / "proj"
    x1 = project / "五点翻倍.xlsx"
    _write_doubled_five_point_xlsx(x1)
    refs = [
        DataRef(title="5点法", relative_path=str(x1.relative_to(project))),
    ]
    data = OriginalRecordData(
        test_item="Dust",
        sample_ids=["A01", "A02", "A03"],
        project_path=str(project),
        data_tables=refs,
        standards=[StdModel(test_name="沙尘")],
    )
    out = tmp_path / "wide.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert "5点法" in paras
    assert "5点法（续）" in paras

    data_tables = []
    for t in doc.tables:
        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is None:
            continue
        widths = [int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))]
        joined = " ".join(
            (_Cell(tc, t).text or "") for row in t.rows for tc in row._tr.tc_lst
        )
        if "9V" in joined and "样品编号 / Sample No." in joined:
            data_tables.append((t, widths, joined))
    assert len(data_tables) == 2
    # First chunk: as many full temp blocks as fit; second starts mid-sequence
    # without breaking a 3-col merge (no orphan 9V without its °C header).
    assert "-40°C" in data_tables[0][2]
    assert "25°C" in data_tables[1][2] or "85°C" in data_tables[1][2]
    for _t, widths, joined in data_tables:
        assert "样品编号 / Sample No." in joined
        assert widths[0] == _OR_SAMPLE_COL_DXA
        assert sum(widths) == _OR_TABLE_WIDTH_DXA
        content = widths[1:]
        assert max(content) - min(content) <= len(content)  # drift on last col
        assert len(content) % 3 == 0
        assert all(w == content[0] for w in content[:-1])


def _png(color: tuple[int, int, int]) -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 20), color).save(buf, format="PNG")
    return buf.getvalue()


def _blip_ids(doc: Document) -> list[str]:
    return [
        blip.get(qn("r:embed"))
        for blip in doc.element.findall(".//" + qn("a:blip"))
        if blip.get(qn("r:embed"))
    ]


def _content_children(doc: Document):
    return [child for child in doc.element.body if child.tag != qn("w:sectPr")]


def _is_bare_page_break(child) -> bool:
    if child.tag != qn("w:p"):
        return False
    breaks = list(child.iter(qn("w:br")))
    return (
        len(breaks) == 1
        and breaks[0].get(qn("w:type")) == "page"
        and not "".join(child.itertext()).strip()
    )


def test_resolve_combined_folder_prefers_remote(tmp_path: Path):
    local = tmp_path / "local"
    remote = tmp_path / "remote"
    local.mkdir()
    remote.mkdir()
    folder = resolve_combined_original_record_folder(local, remote)
    assert folder == remote / TEST_GROUP_DIR
    assert resolve_combined_original_record_folder(local, tmp_path / "missing") == (
        local / TEST_GROUP_DIR
    )
    assert resolve_combined_original_record_folder(None, None) is None


def test_export_leg_original_record_stacks_with_page_break(tmp_path: Path):
    local = tmp_path / "local"
    remote = tmp_path / "remote"
    local.mkdir()
    remote.mkdir()
    heat = OriginalRecordData(
        application_no="A1",
        test_item="高温试验",
        leg_name="Leg 1",
        project_path=str(local),
        remote_root=str(remote),
        standards=[StdModel(test_name="高温试验", standard_desc="高温条件", images=[_png((10, 20, 30))])],
    )
    dust = OriginalRecordData(
        application_no="A1",
        test_item="沙尘试验",
        leg_name="Leg 1",
        project_path=str(local),
        remote_root=str(remote),
        standards=[StdModel(test_name="沙尘试验", standard_desc="沙尘条件", images=[_png((200, 10, 10))])],
    )
    out = export_leg_original_record([heat, dust], template_path=TEMPLATE)
    assert out.parent == remote / TEST_GROUP_DIR
    assert out.name == "A1_Leg_1_原始记录.docx"
    assert not list((remote / TEST_GROUP_DIR).glob("*/*.docx"))

    doc = Document(str(out))
    heat_doc = build_original_record_document(heat, template_path=TEMPLATE)
    dust_doc = build_original_record_document(dust, template_path=TEMPLATE)
    merged = _content_children(doc)
    heat_n = len(_content_children(heat_doc))
    dust_n = len(_content_children(dust_doc))
    assert len(merged) == heat_n + 1 + dust_n
    assert _is_bare_page_break(merged[heat_n])
    assert "".join(merged[heat_n - 1].itertext()) == "".join(
        _content_children(heat_doc)[-1].itertext()
    )
    assert "".join(merged[heat_n + 1].itertext()) == "".join(
        _content_children(dust_doc)[0].itertext()
    )
    before = "".join("".join(child.itertext()) for child in merged[:heat_n])
    after = "".join("".join(child.itertext()) for child in merged[heat_n + 1 :])
    assert "高温试验" in before and "沙尘试验" not in before
    assert "沙尘试验" in after
    embeds = _blip_ids(doc)
    assert len(embeds) >= 2
    blobs = []
    for rid in embeds:
        assert rid in doc.part.rels
        blobs.append(doc.part.rels[rid].target_part.blob)
    assert len(set(blobs)) >= 2

    again = export_leg_original_record([heat], template_path=TEMPLATE)
    assert again.name == "A1_Leg_1_原始记录-2.docx"
