from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from src.generators.elp_docx import clear_field_update_flags, mark_fields_for_update
from src.generators.word_fields import (
    finalize_exported_fields,
    refresh_word_fields,
    skip_word_automation,
    strip_field_update_flags_in_package,
    toc_is_stale,
    toc_result_texts,
)
from src.io.project_mirror import repo_root


TEMPLATE = repo_root() / "templates" / "report_templates" / "template_elp_zh.docx"


def _has_update_fields(path: Path) -> bool:
    doc = Document(str(path))
    flag = doc.settings.element.find(qn("w:updateFields"))
    return flag is not None and flag.get(qn("w:val")) == "true"


def _toc_dirty(path: Path) -> bool:
    doc = Document(str(path))
    for paragraph in doc.element.iter(qn("w:p")):
        instrs = [node.text or "" for node in paragraph.iter(qn("w:instrText"))]
        if not any("TOC" in text.upper() for text in instrs):
            continue
        for fld in paragraph.iter(qn("w:fldChar")):
            if fld.get(qn("w:fldCharType")) == "begin" and fld.get(qn("w:dirty")) == "true":
                return True
    return False


def test_clear_field_update_flags_removes_prompt(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    src = tmp_path / "fields.docx"
    src.write_bytes(TEMPLATE.read_bytes())
    doc = Document(str(src))
    mark_fields_for_update(doc)
    doc.save(str(src))
    assert _has_update_fields(src)
    assert _toc_dirty(src)

    doc = Document(str(src))
    clear_field_update_flags(doc)
    doc.save(str(src))
    assert not _has_update_fields(src)
    assert not _toc_dirty(src)


def test_strip_flags_in_package_keeps_other_parts(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    src = tmp_path / "pack.docx"
    src.write_bytes(TEMPLATE.read_bytes())
    doc = Document(str(src))
    mark_fields_for_update(doc)
    doc.save(str(src))
    before = ZipFile(src).namelist()
    strip_field_update_flags_in_package(src)
    after = ZipFile(src).namelist()
    assert before == after
    assert not _has_update_fields(src)
    assert not _toc_dirty(src)


def test_template_toc_matches_headings():
    if not TEMPLATE.is_file():
        return
    doc = Document(str(TEMPLATE))
    assert toc_result_texts(doc)
    assert toc_is_stale(doc) is False


def test_exported_elp_toc_is_stale_without_word(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from src.generators.elp_engine import ElpReportGenerator
    from tests.test_elp_export import _state

    out = tmp_path / "elp.docx"
    ElpReportGenerator(str(TEMPLATE)).generate(
        _state(tmp_path / "proj"),
        str(out),
        report_no="A226074579110100001E",
        refresh_fields=lambda _path: False,
    )
    doc = Document(str(out))
    assert toc_is_stale(doc) is True
    assert _has_update_fields(out)
    lines = "\n".join(toc_result_texts(doc))
    assert "抛负载" in lines


def test_finalize_success_strips_prompt(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    src = tmp_path / "ok.docx"
    src.write_bytes(TEMPLATE.read_bytes())
    doc = Document(str(src))
    mark_fields_for_update(doc)
    doc.save(str(src))

    refreshed = finalize_exported_fields(src, refresher=lambda _path: True)
    assert refreshed is True
    assert not _has_update_fields(src)
    assert not _toc_dirty(src)


def test_finalize_success_but_stale_falls_back(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from src.generators.elp_engine import ElpReportGenerator
    from tests.test_elp_export import _state

    out = tmp_path / "stale.docx"
    ElpReportGenerator(str(TEMPLATE)).generate(
        _state(tmp_path / "proj"),
        str(out),
        report_no="A226074579110100001E",
        refresh_fields=lambda _path: False,
    )
    # Pretend Word ran but left the cached full TOC.
    refreshed = finalize_exported_fields(out, refresher=lambda _path: True)
    assert refreshed is False
    assert _has_update_fields(out)


def test_refresh_word_fields_skips_under_pytest():
    assert skip_word_automation() is True
    assert refresh_word_fields("/tmp/missing.docx") is False


def test_refresh_word_fields_runs_osascript(monkeypatch, tmp_path: Path):
    doc = tmp_path / "n.docx"
    doc.write_bytes(b"PK\x03\x04fake")
    monkeypatch.setenv("REPORT_CREATOR_REFRESH_WORD", "1")
    monkeypatch.setattr("src.generators.word_fields.sys.platform", "darwin")
    monkeypatch.setattr(
        "src.generators.word_fields._MAC_WORD_APP",
        tmp_path / "Microsoft Word.app",
    )
    (tmp_path / "Microsoft Word.app").mkdir()
    calls = {}

    def fake_run(cmd, **kwargs):
        calls["cmd"] = cmd
        calls["input"] = kwargs.get("input")

        class Result:
            returncode = 0

        return Result()

    monkeypatch.setattr("src.generators.word_fields.subprocess.run", fake_run)
    monkeypatch.setattr("src.generators.word_fields.time.sleep", lambda _s: None)
    assert refresh_word_fields(doc) is True
    assert calls["cmd"][0] == "osascript"
    assert "tables of contents" in calls["input"]
    assert "active document" in calls["input"]
    assert str(doc.resolve()) in calls["cmd"]
