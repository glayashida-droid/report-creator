"""备注/Note storage and where it lands in the three exports."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from PIL import Image

from src.generators.elp_engine import ElpReportGenerator
from src.generators.original_record import OriginalRecordData, generate_original_record
from src.generators.word_engine import WordGenerator
from src.io.test_notes import (
    add_note_images,
    create_note_workbook,
    list_merged_note_table_refs,
    list_note_image_paths,
    retarget_node_note_tables,
)
from src.io.test_photos import list_albums
from src.io.test_photos import test_dir as trial_dir
from src.language_copy import note_section_label
from src.models.project_state import DataTableRef, TestNode
from src.io.project_assets import list_merged_albums

LEG = "Leg 1"
TEST = "色差"
TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "original_data_sheet"
    / "original_record_placeholders.docx"
)


def _png(path: Path, color=(200, 40, 40)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (80, 40), color).save(path, "PNG")
    return path


def test_note_folders_are_not_photo_albums(tmp_path):
    root = tmp_path / "proj"
    folder = trial_dir(root, LEG, TEST)
    (folder / "试验前").mkdir(parents=True)
    (folder / "备注图片").mkdir()
    (folder / "备注表格").mkdir()
    assert list_albums(root, LEG, TEST) == ["试验前"]
    assert list_merged_albums(root, None, LEG, TEST) == ["试验前"]


def test_note_images_keep_order_and_tables_retarget(tmp_path):
    root = tmp_path / "proj"
    first = _png(tmp_path / "a.png")
    second = _png(tmp_path / "b.png", (10, 10, 200))
    add_note_images(root, LEG, TEST, [second, first])
    paths = list_note_image_paths(root, None, LEG, TEST, order=["a.png", "b.png"])
    assert [p.name for p in paths] == ["a.png", "b.png"]

    ref = create_note_workbook(root, LEG, TEST, "灰卡")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "等级"
    ws["B1"] = "5"
    wb.save(root / ref.relative_path)
    wb.close()
    listed = list_merged_note_table_refs(root, None, LEG, TEST)
    assert listed[0].title == "灰卡"
    assert listed[0].relative_path.endswith("备注表格/灰卡.xlsx")

    node = TestNode(test_name=TEST, note_tables=[ref])
    old_key = f"{LEG}-{TEST}"
    new_key = f"{LEG}-色差复测"
    retarget_node_note_tables(node, old_key, new_key)
    assert f"/{new_key}/备注表格/" in node.note_tables[0].relative_path


def test_note_section_labels():
    assert note_section_label("中文") == "备注："
    assert note_section_label("英文") == "Note:"
    assert note_section_label("中英文") == "备注/Note："


def _note_paragraphs(node, root, lang="中文"):
    doc = Document()
    anchor = doc.add_paragraph("ANCHOR")
    engine = WordGenerator("unused.docx")
    engine._report_language = lang
    inserted = engine._insert_note_section(
        doc,
        anchor,
        node,
        str(root),
        LEG,
        label=note_section_label(lang),
    )
    texts = [p.text for p in doc.paragraphs if p.text != "ANCHOR"]
    return inserted, texts, doc


def test_report_note_follows_text_then_image_then_table(tmp_path):
    root = tmp_path / "proj"
    _png(tmp_path / "card.png")
    add_note_images(root, LEG, TEST, [tmp_path / "card.png"])
    ref = create_note_workbook(root, LEG, TEST, "表")
    wb = Workbook()
    wb.active["A1"] = "样品"
    wb.active["B1"] = "结果"
    wb.save(root / ref.relative_path)
    wb.close()
    node = TestNode(
        test_name=TEST,
        note_text="参照 GB/T 250-2008，5 级最好。",
        note_tables=[DataTableRef(title="表", relative_path=ref.relative_path)],
        note_image_order=["card.png"],
    )
    inserted, texts, doc = _note_paragraphs(node, root, "中英文")
    assert inserted
    assert texts[0] == "备注/Note："
    assert texts[1].startswith("参照 GB/T 250-2008")
    assert "试验数据" not in "\n".join(texts)
    assert "表" not in texts
    drawings = list(doc.element.body.iter(qn("w:drawing")))
    assert drawings
    assert any(t.rows[0].cells[0].text == "样品" for t in doc.tables)

    empty = TestNode(test_name="没有备注")
    inserted_empty, texts_empty, _doc = _note_paragraphs(empty, tmp_path / "empty")
    assert not inserted_empty
    assert texts_empty == []


def test_original_record_fills_note_line(tmp_path):
    data = OriginalRecordData(
        application_no="A1",
        test_item=TEST,
        leg_name=LEG,
        project_path=str(tmp_path / "proj"),
        note_text="变色灰卡 5 级最好。",
    )
    out = tmp_path / "record.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))
    blob = "\n".join(p.text for p in doc.paragraphs)
    assert "备注(note)" in blob
    assert "变色灰卡 5 级最好。" in blob
    assert not any(set((p.text or "").strip()) <= {"_"} and (p.text or "").strip() for p in doc.paragraphs)


def test_original_record_keeps_underline_when_empty(tmp_path):
    data = OriginalRecordData(application_no="A1", test_item=TEST, leg_name=LEG)
    out = tmp_path / "record.docx"
    generate_original_record(data, out, template_path=TEMPLATE)
    doc = Document(str(out))
    assert any(set((p.text or "").strip()) <= {"_"} and (p.text or "").strip() for p in doc.paragraphs)


def test_note_panel_switches_text_images_and_tables(tmp_path):
    import sys

    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication(sys.argv)
    root = tmp_path / "proj"
    _png(tmp_path / "keep.png")
    add_note_images(root, LEG, TEST, [tmp_path / "keep.png"])
    node = TestNode(test_name=TEST, note_text="已有备注", note_image_order=["keep.png"])
    from src.ui.test_note_panel import TestNotePanel

    panel = TestNotePanel(node, root, None, LEG)
    panel.show()
    app.processEvents()
    assert panel.stack.currentIndex() == 0
    assert panel.editor.toPlainText() == "已有备注"
    assert panel.editor.height() >= 140
    assert panel.image_order() == ["keep.png"]
    panel.btn_images.click()
    app.processEvents()
    assert panel.stack.currentIndex() == 1
    assert panel.gallery.isVisible()
    panel.btn_tables.click()
    app.processEvents()
    assert panel.stack.currentIndex() == 2
    panel.btn_text.click()
    app.processEvents()
    assert panel.editor.isVisible()
    assert panel.summary_text() == "有文字 · 1 张"
    panel.editor.setPlainText("改过")
    panel.commit_to(node)
    assert node.note_text == "改过"
    assert node.note_image_order == ["keep.png"]
    panel.close()


def test_elp_remark_cell_gets_text_image_and_table(tmp_path):
    root = tmp_path / "proj"
    _png(tmp_path / "card.png")
    add_note_images(root, LEG, TEST, [tmp_path / "card.png"])
    ref = create_note_workbook(root, LEG, TEST, "表")
    wb = Workbook()
    wb.active["A1"] = "等级"
    wb.active["B1"] = "5"
    wb.save(root / ref.relative_path)
    wb.close()
    node = TestNode(
        test_name=TEST,
        note_text="5 级最好",
        note_tables=[ref],
        note_image_order=["card.png"],
    )
    doc = Document()
    table = doc.add_table(1, 2)
    cell = table.rows[0].cells[1]
    gen = ElpReportGenerator("unused.docx")
    gen._fill_remark_cell(
        cell, node, project_path=str(root), remote_root=None, leg_name=LEG
    )
    assert "5 级最好" in cell.text
    assert cell.tables
    assert cell.tables[0].rows[0].cells[0].text == "等级"
    assert list(cell._tc.iter(qn("w:drawing")))

    blank = Document().add_table(1, 1).rows[0].cells[0]
    gen._fill_remark_cell(
        blank,
        TestNode(test_name="没有备注"),
        project_path=str(tmp_path / "empty"),
        remote_root=None,
        leg_name=LEG,
    )
    assert blank.text.strip() == "/"
