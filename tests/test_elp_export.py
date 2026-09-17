from pathlib import Path

from PIL import Image

from src.generators.elp_docx import cell_text, unique_cells
from src.generators.elp_engine import ElpExportFields, ElpReportGenerator
from src.io.network_sources import ELP_REPORT_LANGUAGE, REPORT_LANGUAGE_CHOICES
from src.io.project_mirror import repo_root
from src.io.test_photos import (
    SAMPLE_INFO_DIR_NAME,
    create_album,
    ensure_sample_info_dir,
    list_sample_info_photos,
)
from src.models.project_state import (
    ProjectState,
    TestEquipment,
    TestLeg,
    TestNode,
    TestResult,
    TestSample,
    TestStandard,
)
from src.parsers.elp_plan import ElpPlan


TEMPLATE = repo_root() / "templates" / "report_templates" / "template_elp_zh.docx"
LEG = "Leg 1"
_SLASH_CELLS = {"/", "//", ""}


def _blue_run_texts(doc) -> list[str]:
    from docx.oxml.ns import qn

    from src.generators.elp_docx import BLUE_VALS, _story_elements

    found = []
    for root in _story_elements(doc):
        for run in root.iter(qn("w:r")):
            rPr = run.find(qn("w:rPr"))
            if rPr is None:
                continue
            color = rPr.find(qn("w:color"))
            if color is None:
                continue
            val = (color.get(qn("w:val")) or "").upper()
            if val not in BLUE_VALS:
                continue
            text = "".join((node.text or "") for node in run.findall(qn("w:t")))
            if text.strip():
                found.append(text)
    return found


def _equipment_data(table):
    header = None
    data = []
    for row in table.rows:
        cells = unique_cells(row)
        if len(cells) < 3 or "试验设备" not in cell_text(cells[1]):
            continue
        vals = [cell_text(c) for c in cells]
        if len(cells) >= 6 and vals[2] == "设备名称":
            header = vals
        elif header is not None:
            data.append(vals[2:6] if len(vals) >= 6 else vals[2:])
    return header, data


def _is_slash_row(vals) -> bool:
    return bool(vals) and all((v or "").strip() in _SLASH_CELLS for v in vals)


def _first_detail_table(doc):
    for table in doc.tables:
        header, _ = _equipment_data(table)
        if header is not None:
            return table
    return None


def _png(path: Path, color="red"):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (32, 24), color).save(path, "PNG")
    return path


def test_language_choices_include_elp():
    assert REPORT_LANGUAGE_CHOICES[-1] == ELP_REPORT_LANGUAGE
    assert "中文" in REPORT_LANGUAGE_CHOICES


def test_header_paragraph_uses_center_and_right_tabs():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    paragraph = section.header.paragraphs[0]
    paragraph.text = "试验类型：ECV               报告编号No："
    paragraph.paragraph_format.first_line_indent = 800100
    ElpReportGenerator._rewrite_header_paragraph(
        paragraph, "ELP", "A226074579110100001C", section
    )
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.text == "\t试验类型：ELP\t报告编号No：A226074579110100001C"
    usable = int(ElpReportGenerator._header_usable_width(section))
    stops = list(paragraph.paragraph_format.tab_stops)
    assert [(int(s.position), s.alignment) for s in stops] == [
        (usable // 2, WD_TAB_ALIGNMENT.CENTER),
        (usable, WD_TAB_ALIGNMENT.RIGHT),
    ]
    ind = paragraph._p.find(qn("w:pPr")).find(qn("w:ind")) if paragraph._p.find(qn("w:pPr")) is not None else None
    if ind is not None:
        assert qn("w:firstLine") not in ind.attrib
        assert qn("w:firstLineChars") not in ind.attrib


def test_header_paragraph_report_no_only_is_right_aligned():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    section = doc.sections[0]
    paragraph = section.header.paragraphs[0]
    paragraph.text = "报告编号No："
    ElpReportGenerator._rewrite_header_paragraph(
        paragraph, "ELP", "A226074579110100001C", section
    )
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.text == "报告编号No：A226074579110100001C"
    assert list(paragraph.paragraph_format.tab_stops) == []


def test_cover_value_sits_flush_after_colon():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from src.generators.elp_docx import set_run_color

    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("零件名称：")
    pad = paragraph.add_run("            ")
    pad.underline = True
    blue = paragraph.add_run("XXXXXXX")
    blue.underline = True
    set_run_color(blue, "0000FF")
    tail = paragraph.add_run("                 ")
    tail.underline = True
    ElpReportGenerator._replace_paragraph_blue(paragraph, "副仪表板开关组")
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    after = (paragraph.text or "").split("：", 1)[-1]
    assert after.startswith("副仪表板开关组")


def test_elp_cover_six_lines_are_left_flush(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    gen = ElpReportGenerator(str(TEMPLATE))
    doc = Document(str(TEMPLATE))
    state = _state(tmp_path)
    plan = ElpPlan(product_name="副仪表板开关组", vehicle_code="P166", part_no="6608707284")
    gen._fill_cover(doc, state, state.overview_field_map("中文"), ElpExportFields(), plan)
    expected = {
        "零件名称": "副仪表板开关组",
        "零部件号": "6608707284",
        "项目名称": "P166",
        "试验类型": "ELP",
        "客户名称": "吉利汽车",
        "客户地址": "杭州",
    }
    found = {}
    for paragraph in doc.paragraphs[:20]:
        text = paragraph.text or ""
        for label, value in expected.items():
            if label not in text:
                continue
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
            after = text.split("：", 1)[-1]
            assert after.startswith(value), text
            found[label] = text
    assert set(found) == set(expected)
    company = next(p for p in doc.paragraphs if "上海华测品正" in (p.text or ""))
    assert company.alignment == WD_ALIGN_PARAGRAPH.CENTER
    object_table = next(
        t
        for t in doc.tables
        if "额定电压" in "".join(cell_text(c) for r in t.rows for c in unique_cells(r))
    )
    assert object_table.rows[0].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def _cover_picture_names(paragraph) -> list[str]:
    from docx.oxml.ns import qn

    names = []
    for drawing in paragraph._p.iter(qn("w:drawing")):
        pr = drawing.find(".//" + qn("wp:docPr"))
        if pr is not None and pr.get("name"):
            names.append(pr.get("name"))
    return names


def test_elp_cover_keeps_review_signature(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from docx import Document

    template = Document(str(TEMPLATE))
    address = next(p for p in template.paragraphs[:20] if "客户地址" in (p.text or ""))
    assert "图片 3" in _cover_picture_names(address)

    gen = ElpReportGenerator(str(TEMPLATE))
    doc = Document(str(TEMPLATE))
    state = _state(tmp_path)
    plan = ElpPlan(product_name="副仪表板开关组", vehicle_code="P166", part_no="6608707284")
    gen._fill_cover(doc, state, state.overview_field_map("中文"), ElpExportFields(), plan)
    filled = next(p for p in doc.paragraphs[:20] if "客户地址" in (p.text or ""))
    assert "杭州" in (filled.text or "")
    assert "图片 3" in _cover_picture_names(filled)
    name_line = next(p for p in doc.paragraphs[:20] if "客户名称" in (p.text or ""))
    assert "AutoShape 153" in _cover_picture_names(name_line)


def test_sample_info_dir_created_with_album(tmp_path: Path):
    project = tmp_path / "proj"
    create_album(project, LEG, "启动脉冲", "试验前")
    folder = project / "3.测试组" / SAMPLE_INFO_DIR_NAME
    assert folder.is_dir()
    photo = _png(folder / "正面.png")
    listed = list_sample_info_photos(project)
    assert listed == [photo]
    ensure_sample_info_dir(project)
    assert folder.is_dir()


def test_elp_export_fields_trim_values():
    fields = ElpExportFields(
        rated_voltage=" 12V ",
        sample_source=" 送样 ",
        plan_no="  P166-ELP-TP(副仪表板开关组)-2026-0001  ",
    )
    assert fields.rated_voltage == "12V"
    assert fields.sample_source == "送样"
    assert fields.plan_no == "P166-ELP-TP(副仪表板开关组)-2026-0001"


def _state(project: Path) -> ProjectState:
    node_match = TestNode(
        test_name="启动脉冲",
        standards=[
            TestStandard(
                standard_id="Q/JLY J7111029E-2024",
                chapter="5.2",
                test_name="启动脉冲",
                standard_desc="冷启动/热启动",
                evaluation_req="功能状态A",
                result_desc="功能正常",
            )
        ],
        equipments=[
            TestEquipment(
                name="可编程电源",
                model="NGI",
                code="EQ-1",
                valid_date="2027-01-01",
            )
        ],
        start_date="2026-08-02",
        end_date="2026-08-03",
        env_condition="25℃/50%RH",
        samples=[
            TestSample(sample_id="A01", result=TestResult.PASS, result_desc="功能正常"),
        ],
    )
    node_extra = TestNode(
        test_name="自定义过压",
        standard_id="Q/JLY J7111029E-2024",
        standard_chapter="9.9",
        standard_desc="自定义条件",
        evaluation_req="功能状态A",
        result_desc="功能正常",
        equipment_name="直流电源",
        start_date="2026-08-04",
        end_date="2026-08-05",
        samples=[
            TestSample(sample_id="A01", result=TestResult.PASS, result_desc="功能正常"),
        ],
    )
    return ProjectState(
        project_id="A2260745791101",
        applicant_name="吉利汽车",
        applicant_address="杭州",
        sample_name="副仪表板开关组",
        sample_receive_date="2026-08-01",
        test_start_date="2026-08-02",
        test_end_date="2026-08-10",
        tester_name="测试员",
        application_fields={
            "申请单号": "A22607457911",
            "申请公司": "吉利汽车",
            "申请公司地址": "杭州",
            "样品名称": "副仪表板开关组",
            "零件号": "6608707284",
            "车型代码": "P166",
            "试验类型": "ELP",
            "送样数量": "1",
        },
        legs=[
            TestLeg(leg_id="L1", leg_name=LEG, nodes=[node_match, node_extra]),
        ],
    )


def test_elp_engine_keeps_program_tests_and_fills_fields(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    project = tmp_path / "proj"
    album = create_album(project, LEG, "启动脉冲", "试验前")
    _png(album / "布置.png")
    sample_dir = project / "3.测试组" / SAMPLE_INFO_DIR_NAME
    _png(sample_dir / "正面.png")
    pattern = tmp_path / "patten" / "启动脉冲"
    _png(pattern / "图片10：校准图冷启动t6.png")

    out = tmp_path / "elp.docx"
    plan = ElpPlan(
        product_name="副仪表板开关组",
        vehicle_code="P166",
        part_no="6608707284",
        hw_version="V1.0",
        sw_version="V1.0",
        plan_no="P166-ELP-TP(副仪表板开关组)-2026-0001",
        test_work_modes={"启动脉冲": "工作模式2.2"},
        monitor_rows=[("1", "背光灯点亮", "无闪烁")],
    )
    extras = ElpExportFields(
        rated_voltage="12V",
        sample_source="送样",
        plan_no=plan.plan_no,
    )
    ElpReportGenerator(str(TEMPLATE)).generate(
        _state(project),
        str(out),
        extras=extras,
        plan=plan,
        project_path=str(project),
        pattern_dir=tmp_path / "patten",
        report_no="A226074579110100001E",
    )
    from docx import Document
    from docx.oxml.ns import qn

    from src.generators.elp_docx import iter_body_children, paragraph_style_name, wrap_paragraph

    doc = Document(str(out))
    headings = []
    for child in iter_body_children(doc):
        if child.tag != qn("w:p"):
            continue
        paragraph = wrap_paragraph(child, doc)
        if paragraph_style_name(paragraph).startswith("Heading 2"):
            headings.append((paragraph.text or "").strip())
    assert "启动脉冲" in headings
    assert "自定义过压" in headings
    assert "反极性" not in headings
    assert headings.count("启动脉冲") == 1

    summary = None
    object_table = None
    sample_table = None
    for table in doc.tables:
        blob = "".join(cell_text(c) for r in table.rows for c in unique_cells(r))
        head = "".join(cell_text(c) for r in table.rows[:2] for c in unique_cells(r))
        if "试验项目" in head and "结果判定" in head:
            summary = table
        if "额定电压" in blob and "来样方式" in blob:
            object_table = table
        if "正面" in blob:
            sample_table = table
    assert summary is not None
    names = [cell_text(unique_cells(row)[1]) for row in summary.rows[1:]]
    assert names == ["启动脉冲", "自定义过压"]

    assert object_table is not None
    values = {}
    for row in object_table.rows:
        cells = unique_cells(row)
        if len(cells) >= 2:
            values[cell_text(cells[0])] = cell_text(cells[1])
    assert values["额定电压"] == "12V"
    assert values["来样方式"] == "送样"

    assert sample_table is not None
    captions = [cell_text(c) for c in unique_cells(sample_table.rows[1])]
    assert "正面" in captions

    body = "\n".join(p.text or "" for p in doc.paragraphs)
    tables_text = "\n".join(
        cell_text(c) for t in doc.tables for r in t.rows for c in unique_cells(r)
    )
    assert "P166-ELP-TP(副仪表板开关组)-2026-0001" in tables_text
    assert "7.1.1" in tables_text
    assert "7.2.1" in tables_text
    assert "校准图冷启动t6" in tables_text
    assert "图片10：校准图冷启动t6" not in tables_text
    assert "工作模式2.2" in tables_text
    assert "自定义过压" in (body + tables_text)

    eq_by_test = {}
    for table in doc.tables:
        header, data = _equipment_data(table)
        if header is None:
            continue
        test_name = ""
        for row in table.rows:
            cells = unique_cells(row)
            if len(cells) >= 3 and cell_text(cells[1]) == "试验项目":
                test_name = cell_text(cells[2])
                break
        if test_name:
            eq_by_test[test_name] = data
    assert [row[0] for row in eq_by_test["启动脉冲"]] == ["可编程电源"]
    assert [row[0] for row in eq_by_test["自定义过压"]] == ["直流电源"]
    assert all(not _is_slash_row(row) for rows in eq_by_test.values() for row in rows)

    header = "\n".join(p.text or "" for p in doc.sections[0].header.paragraphs)
    assert "试验类型：ELP" in header.replace(" ", "").replace("\t", "")
    assert "A226074579110100001E" in header
    assert header.count("A226074579110100001E") == 1
    typed = next(p for p in doc.sections[0].header.paragraphs if "试验类型" in (p.text or ""))
    assert "\t试验类型：ELP\t报告编号" in typed.text

    lab_vals = {}
    maker_vals = {}
    for table in doc.tables:
        cells0 = unique_cells(table.rows[0]) if table.rows else []
        if not cells0:
            continue
        if cell_text(cells0[0]) == "名称":
            for row in table.rows:
                pair = unique_cells(row)
                if len(pair) >= 2:
                    lab_vals[cell_text(pair[0])] = cell_text(pair[1])
        if cell_text(cells0[0]) == "制造商":
            for row in table.rows:
                pair = unique_cells(row)
                if len(pair) >= 2:
                    maker_vals[cell_text(pair[0])] = cell_text(pair[1])
    assert lab_vals.get("名称") == "上海华测品正检测技术有限公司"
    assert lab_vals.get("传真") == "/"
    assert maker_vals.get("制造商") == "吉利汽车"
    assert maker_vals.get("电话") == "/"

    settings = doc.settings.element.xml
    assert "updateFields" in settings
    assert _blue_run_texts(doc) == []


def test_elp_engine_pastes_plan_tables(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from docx import Document
    from docx.oxml.ns import qn

    from src.generators.elp_docx import grid_row_groups

    groups = grid_row_groups(
        [
            ["工作模式", "功能描述", "产品功能分类", "", "", "车辆运行模式", "", "", ""],
            ["", "", "A类", "B类", "C类", "Off", "Acc", "Start", "Run"],
            ["mode 1.1", "产品无功能", "√", "", "", "√", "", "", ""],
        ]
    )
    assert groups[0][0] == ("工作模式", 1, False)
    assert groups[0][2] == ("产品功能分类", 3, False)
    assert groups[2][0] == ("mode 1.1", 1, False)
    assert groups[2][2] == ("√", 1, False)

    plan = ElpPlan(
        basic_info_rows=[
            ["电子电器组件种类", "P", "R", "A"],
            ["", "", "", "√"],
            ["工作类型", "连续型", "长时型", "短时型"],
            ["", "√", "", ""],
            ["封装尺寸（不含电缆）", "192.8*58.47*84.2（单位：mm）", "", ""],
        ],
        function_class_rows=[
            ["工作模式", "功能描述", "产品功能分类", "", "", "车辆运行模式", "", "", ""],
            ["", "", "A类", "B类", "C类", "Off", "Acc", "Start", "Run"],
            ["mode 1.1", "产品无功能", "√", "", "", "√", "", "", ""],
            ["Mode 2.2", "产品正常工作", "√", "", "", "", "√", "", ""],
        ],
    )
    out = tmp_path / "elp-plan.docx"
    ElpReportGenerator(str(TEMPLATE)).generate(
        _state(tmp_path / "proj"),
        str(out),
        plan=plan,
        report_no="A226074579110100001C",
    )
    doc = Document(str(out))
    basic = None
    klass = None
    for table in doc.tables:
        head = "".join(cell_text(c) for c in unique_cells(table.rows[0])) if table.rows else ""
        blob = "".join(cell_text(c) for r in table.rows[:3] for c in unique_cells(r))
        if "电子电器组件种类" in head:
            basic = table
        if "产品功能分类" in blob or "产品无功能" in blob:
            klass = table
    assert basic is not None
    basic_text = "\n".join(cell_text(c) for r in basic.rows for c in unique_cells(r))
    assert "封装尺寸" in basic_text
    assert "192.8*58.47*84.2" in basic_text
    ticks = [cell_text(c) for c in unique_cells(basic.rows[1])]
    assert "√" in ticks

    assert klass is not None
    klass_text = "\n".join(cell_text(c) for r in klass.rows for c in unique_cells(r))
    assert "产品无功能" in klass_text
    assert "产品正常工作" in klass_text
    assert "CAN总线通讯" not in klass_text

    header = "".join(p.text or "" for p in doc.sections[0].header.paragraphs)
    assert "\t试验类型：ELP\t报告编号No：" in header
    assert header.index("报告编号") < header.index("A226074579110100001C")
    assert "\nA226074579110100001C" not in "\n".join(
        p.text or "" for p in doc.sections[0].header.paragraphs
    )

    settings_el = doc.settings.element
    flag = settings_el.find(qn("w:updateFields"))
    assert flag is not None
    assert flag.get(qn("w:val")) == "true"
    assert _blue_run_texts(doc) == []


def test_elp_equipment_rows_match_actual_count():
    if not TEMPLATE.is_file():
        return
    from docx import Document

    gen = ElpReportGenerator(str(TEMPLATE))

    def fill(items=None, equipment_name=""):
        doc = Document(str(TEMPLATE))
        table = _first_detail_table(doc)
        node = TestNode(
            test_name="t",
            equipments=list(items or []),
            equipment_name=equipment_name,
        )
        gen._fill_equipment_rows(table, node)
        return table

    def labels(table):
        out = []
        for row in table.rows:
            cells = unique_cells(row)
            if len(cells) >= 2:
                out.append(cell_text(cells[1]))
        return out

    empty = fill([])
    header, data = _equipment_data(empty)
    assert header is not None
    assert header[2:6] == ["设备名称", "设备型号", "内部编号", "校准有效期"]
    assert data == []
    labs = labels(empty)
    assert labs.index("试验室温湿度") == labs.index("试验设备信息") + 1

    two = [
        TestEquipment(name="A", model="m1", code="c1", valid_date="2027-01-01"),
        TestEquipment(name="B", model="m2", code="c2", valid_date="2027-02-01"),
    ]
    filled = fill(two)
    _, data = _equipment_data(filled)
    assert data == [
        ["A", "m1", "c1", "2027-01-01"],
        ["B", "m2", "c2", "2027-02-01"],
    ]
    assert not any(_is_slash_row(row) for row in data)
    labs = labels(filled)
    assert labs.count("试验设备信息") == 3
    assert labs.index("试验室温湿度") == max(
        i for i, lab in enumerate(labs) if lab == "试验设备信息"
    ) + 1

    many = [
        TestEquipment(name=f"E{i}", model="m", code=f"C{i}", valid_date="2027-01-01")
        for i in range(12)
    ]
    overflow = fill(many)
    _, data = _equipment_data(overflow)
    assert [row[0] for row in data] == [f"E{i}" for i in range(12)]
    labs = labels(overflow)
    assert labs.index("试验室温湿度") == max(
        i for i, lab in enumerate(labs) if lab == "试验设备信息"
    ) + 1


def _block_live_tables(doc, block):
    from docx.oxml.ns import qn

    from src.generators.elp_docx import wrap_table

    return [
        wrap_table(el, doc)
        for el in block.elements
        if el.tag == qn("w:tbl") and el.getparent() is not None
    ]


def _fill_first_block_photos(doc, paths):
    from docx.oxml.ns import qn

    from src.generators.elp_docx import wrap_table

    gen = ElpReportGenerator(str(TEMPLATE))
    block = gen._collect_test_blocks(doc)[0]
    tables = [wrap_table(el, doc) for el in block.elements if el.tag == qn("w:tbl")]
    gen._collect_photos = lambda *a, **k: list(paths)
    gen._fill_photo_tables(
        tables[1:],
        TestNode(test_name="自定义试验"),
        doc=doc,
        leg_name=LEG,
        project_path=None,
        remote_root=None,
        pattern_dir=None,
        matched=False,
    )
    return block


def test_elp_photo_tables_drop_empty_slots(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from docx import Document

    photos = [_png(tmp_path / f"校准{i}.png") for i in range(1, 6)]
    doc = Document(str(TEMPLATE))
    block = _fill_first_block_photos(doc, photos)
    photo_tables = _block_live_tables(doc, block)[1:]
    assert len(photo_tables) == 1
    table = photo_tables[0]
    assert len(table.rows) == 7
    last_img = unique_cells(table.rows[5])
    last_cap = unique_cells(table.rows[6])
    assert len(last_img) == 1
    assert len(last_cap) == 1
    prev_img = unique_cells(table.rows[3])
    assert len(prev_img) == 2
    captions = [
        cell_text(c)
        for row_i in (2, 4, 6)
        for c in unique_cells(table.rows[row_i])
    ]
    assert captions == [
        "图片1：校准1",
        "图片2：校准2",
        "图片3：校准3",
        "图片4：校准4",
        "图片5：校准5",
    ]
    from docx.oxml.ns import qn

    last_tcPr = last_img[0]._tc.find(qn("w:tcPr"))
    last_w = last_tcPr.find(qn("w:tcW"))
    prev_tcPr = prev_img[0]._tc.find(qn("w:tcPr"))
    prev_w = prev_tcPr.find(qn("w:tcW"))
    assert last_w is not None and prev_w is not None
    assert last_w.get(qn("w:w")) == prev_w.get(qn("w:w"))
    assert last_tcPr.find(qn("w:gridSpan")) is None
    blob = "\n".join(
        cell_text(c) for t in photo_tables for r in t.rows for c in unique_cells(r)
    )
    assert "图片6" not in blob
    assert "XXXXXXX" not in blob


def test_elp_photo_tables_removed_when_empty():
    if not TEMPLATE.is_file():
        return
    from docx import Document

    doc = Document(str(TEMPLATE))
    block = _fill_first_block_photos(doc, [])
    live = _block_live_tables(doc, block)
    assert len(live) == 1
    assert "相关图片" not in "".join(
        cell_text(c) for t in live for r in t.rows for c in unique_cells(r)
    )


def test_elp_slim_block_keeps_one_photo_table():
    if not TEMPLATE.is_file():
        return
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(TEMPLATE))
    gen = ElpReportGenerator(str(TEMPLATE))
    block = gen._collect_test_blocks(doc)[0]
    slim = gen._slim_block_elements(block.elements)
    assert sum(1 for el in slim if el.tag == qn("w:tbl")) == 2
    assert sum(1 for el in block.elements if el.tag == qn("w:tbl")) == 5


def _tables_until_next_heading(doc, heading_el):
    from docx.oxml.ns import qn

    from src.generators.elp_docx import paragraph_style_name, wrap_paragraph, wrap_table

    tables = []
    el = heading_el.getnext()
    while el is not None:
        if el.tag == qn("w:p"):
            para = wrap_paragraph(el, doc)
            if paragraph_style_name(para).startswith("Heading 2"):
                break
        if el.tag == qn("w:tbl"):
            tables.append(wrap_table(el, doc))
        el = el.getnext()
    return tables


def test_elp_photo_tables_grow_when_needed(tmp_path: Path):
    if not TEMPLATE.is_file():
        return
    from docx import Document
    from docx.oxml.ns import qn

    from src.generators.elp_docx import delete_table, wrap_table

    photos = [_png(tmp_path / f"校准{i}.png") for i in range(1, 10)]
    doc = Document(str(TEMPLATE))
    gen = ElpReportGenerator(str(TEMPLATE))
    block = gen._collect_test_blocks(doc)[0]
    tables = [wrap_table(el, doc) for el in block.elements if el.tag == qn("w:tbl")]
    for extra in tables[2:]:
        delete_table(extra)
    gen._collect_photos = lambda *a, **k: list(photos)
    gen._fill_photo_tables(
        tables[1:2],
        TestNode(test_name="自定义试验"),
        doc=doc,
        leg_name=LEG,
        project_path=None,
        remote_root=None,
        pattern_dir=None,
        matched=False,
    )
    live = _tables_until_next_heading(doc, block.heading_el)
    assert len(live) == 3
    captions = [
        cell_text(c)
        for table in live[1:]
        for row_i, row in enumerate(table.rows)
        if row_i in (2, 4, 6, 8)
        for c in unique_cells(row)
        if cell_text(c)
    ]
    assert captions[0] == "图片1：校准1"
    assert captions[-1] == "图片9：校准9"
    assert len(captions) == 9


def test_elp_template_keeps_blue_slots():
    if not TEMPLATE.is_file():
        return
    from docx import Document

    assert _blue_run_texts(Document(str(TEMPLATE)))


def test_set_blue_portion_writes_black():
    from docx import Document

    from src.generators.elp_docx import (
        BLACK_VAL,
        cell_text,
        run_color_val,
        run_is_blue,
        set_blue_portion,
        set_run_color,
    )

    doc = Document()
    cell = doc.add_table(1, 1).cell(0, 0)
    run = cell.paragraphs[0].add_run("slot")
    set_run_color(run, "0000FF")
    assert run_is_blue(run)
    assert set_blue_portion(cell, "filled")
    assert cell_text(cell) == "filled"
    filled = cell.paragraphs[0].runs[0]
    assert not run_is_blue(filled)
    assert run_color_val(filled) == BLACK_VAL


def test_paint_blue_runs_black_covers_header():
    from docx import Document

    from src.generators.elp_docx import paint_blue_runs_black, set_run_color

    doc = Document()
    run = doc.sections[0].header.paragraphs[0].add_run("ELP")
    set_run_color(run, "0000FF")
    assert _blue_run_texts(doc) == ["ELP"]
    paint_blue_runs_black(doc)
    assert _blue_run_texts(doc) == []
