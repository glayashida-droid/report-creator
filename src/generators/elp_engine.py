"""Fill the Geely ELP Word template from project state + ELP test plan."""

from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt
from docx.table import Table

from src.generators.elp_docx import (
    BLACK_VAL,
    cell_has_blue,
    cell_text,
    delete_row_cell,
    delete_table,
    delete_table_row,
    fill_result_slot,
    force_black_text,
    force_blue_text,
    insert_cloned_row_after,
    insert_cloned_table_after,
    iter_body_children,
    paint_blue_runs_black,
    paragraph_style_name,
    remove_drawings,
    run_has_drawing,
    set_blue_portion,
    set_photo_caption,
    set_run_color,
    unique_cells,
    wrap_paragraph,
    wrap_table,
    write_runs,
)
from src.generators.word_engine import PHOTO_WIDTH_IN, WordGenerator
from src.generators.word_fields import finalize_exported_fields
from src.io.network_sources import elp_data_pattern_directory
from src.io.test_photos import is_image_file, list_sample_info_photos
from src.models.project_state import ProjectState, TestEquipment, TestNode
from src.parsers.elp_plan import (
    CHAPTER6_BASIC,
    CHAPTER6_FUNCTION_CLASS,
    CHAPTER6_MONITOR,
    CHAPTER6_STATES,
    CHAPTER6_WORK_MODES,
    ElpPlan,
    caption_from_stem,
    elp_names_match,
    lookup_work_mode,
)

LAB_NAME_CN = "上海华测品正检测技术有限公司"
LAB_PHONE = "021-31073300"
LAB_ZIP = "201114"
ELP_LAB_ADDRESS_CN = "上海市闵行区新骏环路 777 号"
DIFF_DEFAULT = "测试计划中的测试方式及相关条件与标准无差异"


class ElpExportFields:
    def __init__(
        self,
        rated_voltage: str = "",
        sample_source: str = "",
        plan_no: str = "",
        manufacturer_address: str = "",
    ):
        self.rated_voltage = (rated_voltage or "").strip()
        self.sample_source = (sample_source or "").strip()
        self.plan_no = (plan_no or "").strip()
        self.manufacturer_address = (manufacturer_address or "").strip()


def manufacturer_address_prompt(fields: Dict[str, str], state: ProjectState) -> Optional[str]:
    """Applicant address to confirm, or None when the form already has 生产商地址."""
    if (fields.get("生产商地址") or "").strip():
        return None
    return (fields.get("申请公司地址") or state.applicant_address or "").strip()


class _TestBlock:
    def __init__(self, elements: List, heading_el, detail_tbl_el):
        self.elements = elements
        self.heading_el = heading_el
        self.detail_tbl_el = detail_tbl_el
        self.test_name = ""


class ElpReportGenerator:
    def __init__(self, template_path: str):
        self.template_path = template_path
        self._word = WordGenerator(template_path)
        self._word._report_language = "中文"
        self._export_temps: list = []
        self.toc_refreshed = False

    def generate(
        self,
        state: ProjectState,
        output_path: str,
        *,
        project_path: Optional[str] = None,
        remote_root: Optional[str] = None,
        leg_filter: Optional[str] = None,
        report_no: str = "",
        extras: Optional[ElpExportFields] = None,
        plan: Optional[ElpPlan] = None,
        pattern_dir: Optional[Path] = None,
        refresh_fields: Optional[Callable[[str], bool]] = None,
    ) -> None:
        extras = extras or ElpExportFields()
        plan = plan or ElpPlan()
        self._export_temps = []
        self.toc_refreshed = False
        try:
            doc = Document(self.template_path)
            pairs = list(state.iter_nodes_for_export(leg_filter))
            nodes = [node for _leg, node in pairs]
            fields = state.overview_field_map("中文")
            self._fill_cover(doc, state, fields, extras, plan)
            self._fill_headers(doc, fields.get("试验类型") or "", report_no)
            self._fill_object_table(doc, state, fields, extras, plan, nodes)
            self._fill_lab_and_maker(doc, state, fields, extras)
            self._prefetch_export_photos(
                pairs,
                project_path,
                remote_root,
                pattern_dir,
            )
            self._fill_sample_photos(doc, project_path)
            self._fill_plan_chapter6(doc, plan, nodes)
            self._rewrite_chapter7(
                doc,
                state,
                pairs,
                fields,
                extras,
                plan,
                project_path,
                remote_root,
                pattern_dir,
            )
            self._fill_summary_table(doc, nodes)
            self._fill_diff_table(doc, nodes)
            paint_blue_runs_black(doc)
            doc.save(output_path)
            self.toc_refreshed = finalize_exported_fields(
                output_path, refresher=refresh_fields
            )
        finally:
            for temp in self._export_temps:
                try:
                    Path(temp).unlink(missing_ok=True)
                except OSError:
                    pass
            self._export_temps = []

    # ------------------------------------------------------------------ cover / headers

    def _fill_cover(
        self,
        doc: Document,
        state: ProjectState,
        fields: Dict[str, str],
        extras: ElpExportFields,
        plan: ElpPlan,
    ) -> None:
        sample = fields.get("样品名称") or state.sample_name or plan.product_name
        part = fields.get("零件号") or plan.part_no
        project = fields.get("车型代码") or plan.vehicle_code
        test_type = fields.get("试验类型") or ""
        customer = fields.get("申请公司") or state.applicant_name
        address = fields.get("申请公司地址") or state.applicant_address
        mapping = {
            "零件名称": sample,
            "零部件号": part,
            "项目名称": project,
            "试验类型": test_type,
            "客户名称": customer,
            "客户地址": address,
        }
        for paragraph in doc.paragraphs[:20]:
            text = paragraph.text or ""
            for label, value in mapping.items():
                if label in text:
                    self._replace_paragraph_blue(paragraph, value)
                    break

    @staticmethod
    def _replace_paragraph_blue(paragraph, value: str) -> None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        runs = list(paragraph.runs)
        blue = [run for run in runs if _run_is_blue(run)]
        if not blue:
            return
        for run in runs:
            if run in blue:
                break
            if run_has_drawing(run):
                continue
            if not (run.text or "").strip():
                run.text = ""
        blue[0].text = value or ""
        set_run_color(blue[0], BLACK_VAL)
        for run in blue[1:]:
            run.text = ""
            set_run_color(run, BLACK_VAL)

    def _fill_headers(self, doc: Document, test_type: str, report_no: str) -> None:
        for section in doc.sections:
            parts = [section.header]
            try:
                if section.different_first_page_header_footer:
                    parts.append(section.first_page_header)
            except Exception:
                pass
            for part in parts:
                if part is None:
                    continue
                for paragraph in part.paragraphs:
                    self._rewrite_header_paragraph(paragraph, test_type, report_no, section)
                    text = (paragraph.text or "").strip()
                    if "试验类型" in text or "报告编号" in text:
                        continue
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)

    @staticmethod
    def _header_usable_width(section):
        page = section.page_width
        left = section.left_margin or 0
        right = section.right_margin or 0
        return page - left - right

    @staticmethod
    def _clear_header_first_indent(paragraph) -> None:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        pPr = paragraph._p.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            for attr in (qn("w:firstLine"), qn("w:firstLineChars")):
                if attr in ind.attrib:
                    del ind.attrib[attr]

    @staticmethod
    def _set_header_tab_stops(paragraph, section, *, center: bool, right: bool) -> None:
        stops = paragraph.paragraph_format.tab_stops
        stops.clear_all()
        usable = int(ElpReportGenerator._header_usable_width(section))
        if center:
            stops.add_tab_stop(Emu(usable // 2), WD_TAB_ALIGNMENT.CENTER)
        if right:
            stops.add_tab_stop(Emu(usable), WD_TAB_ALIGNMENT.RIGHT)

    @staticmethod
    def _rewrite_header_paragraph(paragraph, test_type: str, report_no: str, section) -> None:
        text = paragraph.text or ""
        has_type = "试验类型" in text
        has_no = "报告编号" in text
        if not has_type and not has_no:
            return
        no_label = ""
        if has_no:
            if "No." in text:
                no_label = "报告编号No.："
            elif "No" in text:
                no_label = "报告编号No："
            else:
                no_label = "报告编号："
        pieces: List[Tuple[str, Optional[str]]] = []
        if has_type and has_no:
            pieces.extend(
                [
                    ("\t", None),
                    ("试验类型：", None),
                    (test_type or "", None),
                    ("\t", None),
                    (no_label, None),
                    (report_no or "", None),
                ]
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ElpReportGenerator._set_header_tab_stops(paragraph, section, center=True, right=True)
        elif has_type:
            pieces.append(("试验类型：", None))
            pieces.append((test_type or "", None))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.tab_stops.clear_all()
        else:
            pieces.append((no_label, None))
            pieces.append((report_no or "", None))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.tab_stops.clear_all()
        write_runs(paragraph, pieces)
        ElpReportGenerator._clear_header_first_indent(paragraph)

    # ------------------------------------------------------------------ early tables

    def _table_by_labels(self, doc: Document, *needles: str) -> Optional[Table]:
        for table in doc.tables:
            blob = "".join(cell_text(cell) for row in table.rows[:2] for cell in unique_cells(row))
            if all(n in blob for n in needles):
                return table
        return None

    def _table_by_first_label(self, doc: Document, label: str) -> Optional[Table]:
        for table in doc.tables:
            if not table.rows:
                continue
            cells = unique_cells(table.rows[0])
            if cells and cell_text(cells[0]) == label:
                return table
        return None

    def _table_with_first_row(self, doc: Document, needle: str) -> Optional[Table]:
        for table in doc.tables:
            if not table.rows:
                continue
            blob = "".join(cell_text(cell) for cell in unique_cells(table.rows[0]))
            if needle in blob:
                return table
        return None

    def _fill_object_table(
        self,
        doc: Document,
        state: ProjectState,
        fields: Dict[str, str],
        extras: ElpExportFields,
        plan: ElpPlan,
        nodes: Sequence[TestNode],
    ) -> None:
        table = self._table_by_labels(doc, "车型代码", "样品名称")
        if table is None:
            return
        ids = WordGenerator._collect_sample_ids(nodes)
        id_text = WordGenerator._format_id_range(ids) if ids else "/"
        qty = (fields.get("送样数量") or "").strip()
        if not qty:
            qty = str(max((len([s for s in n.samples if s.sample_id]) for n in nodes), default=0) or "")
        values = {
            "车型代码": fields.get("车型代码") or plan.vehicle_code,
            "样品名称": fields.get("样品名称") or state.sample_name or plan.product_name,
            "零部件号": fields.get("零件号") or plan.part_no,
            "软件版本": plan.sw_version,
            "硬件版本": plan.hw_version,
            "样品数量": qty,
            "额定电压": extras.rated_voltage,
            "样品编号": id_text,
            "来样方式": extras.sample_source,
            "样品接收日期": WordGenerator._fmt_date(state.sample_receive_date),
            "测试开始时间": WordGenerator._fmt_date(state.test_start_date),
            "测试结束时间": WordGenerator._fmt_date(state.test_end_date),
        }
        for row in table.rows:
            cells = unique_cells(row)
            if len(cells) < 2:
                continue
            label = cell_text(cells[0])
            for key, val in values.items():
                if label.startswith(key):
                    self._write_blue(cells[1], val)
                    break

    def _fill_lab_and_maker(
        self,
        doc: Document,
        state: ProjectState,
        fields: Dict[str, str],
        extras: ElpExportFields,
    ) -> None:
        lab = self._table_by_first_label(doc, "名称")
        if lab is not None:
            mapping = {
                "名称": LAB_NAME_CN,
                "地址": ELP_LAB_ADDRESS_CN,
                "电话": LAB_PHONE,
                "传真": "/",
                "邮编": LAB_ZIP,
                "E-mail": "/",
            }
            self._fill_kv_table(lab, mapping)
        maker = self._table_by_first_label(doc, "制造商")
        if maker is not None:
            producer = fields.get("生产商") or fields.get("申请公司") or state.applicant_name
            parsed = (fields.get("生产商地址") or "").strip()
            addr = parsed or (extras.manufacturer_address or "").strip()
            self._fill_kv_table(
                maker,
                {
                    "制造商": producer,
                    "地址": addr,
                    "电话": "/",
                    "传真": "/",
                    "邮编": "/",
                },
            )

    def _fill_kv_table(self, table: Table, mapping: Dict[str, str]) -> None:
        for row in table.rows:
            cells = unique_cells(row)
            if len(cells) < 2:
                continue
            label = cell_text(cells[0])
            for key, val in mapping.items():
                if label.startswith(key):
                    self._write_blue(cells[1], val)
                    break

    def _fill_sample_photos(self, doc: Document, project_path: Optional[str]) -> None:
        table = None
        for candidate in doc.tables:
            blob = "".join(
                cell_text(cell) for row in candidate.rows for cell in unique_cells(row)
            )
            if "测试样品照" in blob:
                table = candidate
                break
        if table is None:
            return
        photos: List[Path] = []
        if project_path:
            photos = list_sample_info_photos(Path(project_path))
        slots = [
            (1, 0),
            (1, 1),
            (3, 0),
            (3, 1),
        ]
        image_rows = {1: 0, 3: 2}
        for i, (cap_r, cap_c) in enumerate(slots):
            if cap_r >= len(table.rows):
                continue
            cap_cells = unique_cells(table.rows[cap_r])
            img_r = image_rows[cap_r]
            if img_r >= len(table.rows) or cap_c >= len(cap_cells):
                continue
            img_cells = unique_cells(table.rows[img_r])
            cap_cell = cap_cells[cap_c]
            img_cell = img_cells[cap_c] if cap_c < len(img_cells) else cap_cell
            remove_drawings(img_cell)
            if i < len(photos):
                path = photos[i]
                force_blue_text(cap_cell, path.stem)
                self._put_picture(img_cell, path)
            else:
                force_blue_text(cap_cell, "")

    # ------------------------------------------------------------------ chapter 6

    def _fill_plan_chapter6(
        self, doc: Document, plan: ElpPlan, nodes: Sequence[TestNode]
    ) -> None:
        images = plan.chapter6_images or {}
        if not images:
            return
        modes = None
        klass = None
        for table in list(doc.tables):
            head = cell_text(unique_cells(table.rows[0])[0]) if table.rows else ""
            blob = "".join(cell_text(c) for r in table.rows[:2] for c in unique_cells(r))
            if head == "工作模式" and "定义描述" in blob and modes is None:
                modes = table
            elif "产品功能分类" in blob and klass is None:
                klass = table
        slots = (
            (CHAPTER6_BASIC, self._table_with_first_row(doc, "电子电器组件种类")),
            (CHAPTER6_FUNCTION_CLASS, klass),
            (CHAPTER6_WORK_MODES, modes),
            (CHAPTER6_MONITOR, self._table_by_labels(doc, "功能描述", "可接受范围")),
            (CHAPTER6_STATES, self._table_by_labels(doc, "功能状态", "定义描述")),
        )
        for key, table in slots:
            blob = images.get(key)
            if table is not None and blob:
                self._replace_table_with_picture(doc, table, blob)

    def _replace_table_with_picture(self, doc: Document, table: Table, blob: bytes) -> None:
        new_p = OxmlElement("w:p")
        table._tbl.addnext(new_p)
        paragraph = wrap_paragraph(new_p, doc)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run()
        run.add_picture(io.BytesIO(blob), width=Emu(self._body_content_width(doc)))
        delete_table(table)

    @staticmethod
    def _body_content_width(doc: Document) -> int:
        section = doc.sections[0]
        width = int(section.page_width - section.left_margin - section.right_margin)
        return width if width > 0 else int(Inches(6.3))

    def _fill_summary_table(self, doc: Document, nodes: Sequence[TestNode]) -> None:
        table = self._table_by_labels(doc, "试验项目", "试验标准", "结果判定")
        if table is None:
            return
        originals = []
        for row in table.rows[1:]:
            cells = unique_cells(row)
            if len(cells) < 4:
                continue
            originals.append(
                (cell_text(cells[1]), cell_text(cells[2]), cells)
            )
        template_tr = deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
        while len(table.rows) > 1:
            delete_table_row(table, 1)
        if template_tr is None:
            return
        for idx, node in enumerate(nodes):
            table._tbl.append(deepcopy(template_tr))
            cells = unique_cells(table.rows[-1])
            name = (node.test_name or "").strip()
            std = ""
            matched = next((item for item in originals if elp_names_match(item[0], name)), None)
            if matched:
                name = matched[0]
                std = matched[1]
            else:
                std = WordGenerator._format_method(node) or "/"
            force_black_text(cells[0], str(idx + 1))
            force_black_text(cells[1], name)
            force_black_text(cells[2], std)
            force_blue_text(cells[3], self._conclusion(node))

    def _fill_diff_table(self, doc: Document, nodes: Sequence[TestNode]) -> None:
        table = self._table_by_labels(doc, "测试项目", "测试方法差异说明")
        if table is None:
            return
        originals = []
        for row in table.rows[1:]:
            cells = unique_cells(row)
            if len(cells) < 2:
                continue
            originals.append(cell_text(cells[0]))
        template_tr = deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
        while len(table.rows) > 1:
            delete_table_row(table, 1)
        if template_tr is None:
            return
        for node in nodes:
            table._tbl.append(deepcopy(template_tr))
            cells = unique_cells(table.rows[-1])
            name = (node.test_name or "").strip()
            matched = next((item for item in originals if elp_names_match(item, name)), None)
            force_black_text(cells[0], matched or name)
            force_blue_text(cells[1], DIFF_DEFAULT)

    # ------------------------------------------------------------------ chapter 7

    def _rewrite_chapter7(
        self,
        doc: Document,
        state: ProjectState,
        pairs: Sequence[Tuple],
        fields: Dict[str, str],
        extras: ElpExportFields,
        plan: ElpPlan,
        project_path: Optional[str],
        remote_root: Optional[str],
        pattern_dir: Optional[Path],
    ) -> None:
        blocks = self._collect_test_blocks(doc)
        if not blocks:
            return
        prototype = self._slim_block_elements(deepcopy(blocks[0].elements))
        used = set()
        ordered: List[Tuple[TestNode, List, bool, str]] = []
        for leg, node in pairs:
            match = None
            for idx, block in enumerate(blocks):
                if idx in used:
                    continue
                if elp_names_match(node.test_name, block.test_name):
                    match = block
                    used.add(idx)
                    break
            if match is not None:
                ordered.append((node, match.elements, True, leg.leg_name))
            else:
                ordered.append(
                    (node, [deepcopy(el) for el in prototype], False, leg.leg_name)
                )
        self._remove_elements(doc, [el for block in blocks for el in block.elements])
        anchor = self._chapter7_anchor(doc)
        parent = doc.element.body
        insert_at = list(parent).index(anchor) + 1 if anchor is not None else len(list(parent)) - 1
        for idx, (node, elements, matched, leg_name) in enumerate(ordered):
            for el in elements:
                parent.insert(insert_at, el)
                insert_at += 1
            self._fill_test_block(
                doc,
                elements,
                node,
                index=idx + 1,
                matched=matched,
                leg_name=leg_name,
                state=state,
                fields=fields,
                extras=extras,
                plan=plan,
                project_path=project_path,
                remote_root=remote_root,
                pattern_dir=pattern_dir,
            )

    def _chapter7_anchor(self, doc):
        for child in iter_body_children(doc):
            if child.tag != qn("w:p"):
                continue
            paragraph = wrap_paragraph(child, doc)
            if paragraph_style_name(paragraph).startswith("Heading 1") and "检测项目" in (
                paragraph.text or ""
            ):
                return child
        return None

    def _collect_test_blocks(self, doc: Document) -> List[_TestBlock]:
        started = False
        blocks: List[_TestBlock] = []
        current: List = []
        heading_el = None
        detail_el = None
        for child in iter_body_children(doc):
            if child.tag == qn("w:p"):
                paragraph = wrap_paragraph(child, doc)
                style = paragraph_style_name(paragraph)
                text = (paragraph.text or "").strip()
                if style.startswith("Heading 1") and "检测项目" in text:
                    started = True
                    continue
                if not started:
                    continue
                if style.startswith("Heading 2"):
                    if current and heading_el is not None:
                        block = _TestBlock(current, heading_el, detail_el)
                        block.test_name = self._block_name(doc, block)
                        blocks.append(block)
                    current = [child]
                    heading_el = child
                    detail_el = None
                    continue
            if started and current:
                current.append(child)
                if child.tag == qn("w:tbl") and detail_el is None:
                    detail_el = child
        if started and current and heading_el is not None:
            block = _TestBlock(current, heading_el, detail_el)
            block.test_name = self._block_name(doc, block)
            blocks.append(block)
        return blocks

    @staticmethod
    def _slim_block_elements(elements: Sequence) -> List:
        """Heading + detail table + one photo table (and its spacers)."""
        slim: List = []
        tables = 0
        for el in elements:
            if el.tag == qn("w:tbl"):
                tables += 1
                if tables > 2:
                    continue
                slim.append(el)
                continue
            if tables > 2:
                continue
            slim.append(el)
        return slim

    def _block_name(self, doc: Document, block: _TestBlock) -> str:
        if block.heading_el is not None:
            name = (wrap_paragraph(block.heading_el, doc).text or "").strip()
            if name:
                return name
        if block.detail_tbl_el is None:
            return ""
        table = wrap_table(block.detail_tbl_el, doc)
        for row in table.rows:
            cells = unique_cells(row)
            if len(cells) >= 3 and "试验项目" in cell_text(cells[1]):
                return cell_text(cells[2])
        return ""

    @staticmethod
    def _remove_elements(doc: Document, elements: Sequence) -> None:
        body = doc.element.body
        for el in elements:
            parent = el.getparent()
            if parent is body:
                body.remove(el)

    def _fill_test_block(
        self,
        doc: Document,
        elements: List,
        node: TestNode,
        *,
        index: int,
        matched: bool,
        leg_name: str,
        state: ProjectState,
        fields: Dict[str, str],
        extras: ElpExportFields,
        plan: ElpPlan,
        project_path: Optional[str],
        remote_root: Optional[str],
        pattern_dir: Optional[Path],
    ) -> None:
        tables = [wrap_table(el, doc) for el in elements if el.tag == qn("w:tbl")]
        if not tables:
            return
        detail = tables[0]
        self._renumber_detail(detail, index)
        sample_name = fields.get("样品名称") or state.sample_name or ""
        qty = str(len([s for s in (node.samples or []) if s.sample_id]) or "")
        plan_no = extras.plan_no or plan.plan_no
        mode = lookup_work_mode(plan.test_work_modes, node.test_name)
        env = node.resolved_env_condition() or "/"
        period = self._period(node.start_date, node.end_date)
        tester = (state.tester_name or "").strip() or "/"
        result_text = self._result_text(node)
        conclusion = self._conclusion(node)
        method = WordGenerator._format_method(node) or "/"
        standard = method
        eval_req = (node.evaluation_req or node.joined_evaluation_req() or "").strip()
        for row in detail.rows:
            cells = unique_cells(row)
            if len(cells) < 3:
                continue
            label = cell_text(cells[1])
            value_cell = cells[2]
            if "样品名称" in label:
                self._write_blue(value_cell, sample_name)
            elif "试验项目" in label:
                if not matched:
                    force_black_text(value_cell, node.test_name or "")
            elif "试验标准" in label:
                if not matched:
                    force_black_text(value_cell, standard)
            elif "工作模式" in label:
                if mode:
                    self._write_blue(value_cell, mode)
            elif "样品" in label and "数量" in label:
                self._write_blue(value_cell, qty)
            elif "试验地点" in label:
                self._write_blue(value_cell, LAB_NAME_CN)
            elif "试验室温湿度" in label:
                self._write_blue(value_cell, env)
            elif "试验人员" in label:
                self._write_blue(value_cell, tester)
            elif "试验起止日期" in label:
                self._write_blue(value_cell, period)
            elif "测试计划编号" in label:
                self._write_blue(value_cell, plan_no)
            elif "试验方法" in label:
                if not matched:
                    force_black_text(value_cell, method)
            elif "试验要求" in label:
                if eval_req and (cell_has_blue(value_cell) or not matched):
                    if matched:
                        set_blue_portion(value_cell, eval_req)
                    else:
                        force_black_text(value_cell, eval_req)
            elif "试验结果" in label:
                fill_result_slot(value_cell, result_text)
            elif "结果判定" in label:
                self._write_blue(value_cell, conclusion)
            elif label.startswith("备注"):
                self._fill_remark_cell(
                    value_cell,
                    node,
                    project_path=project_path,
                    remote_root=remote_root,
                    leg_name=leg_name,
                )
        self._fill_equipment_rows(detail, node)
        heading = next((el for el in elements if el.tag == qn("w:p")), None)
        if heading is not None and not matched:
            paragraph = wrap_paragraph(heading, doc)
            if paragraph.runs:
                paragraph.runs[0].text = node.test_name or ""
                for run in paragraph.runs[1:]:
                    run.text = ""
        photo_tables = tables[1:]
        self._fill_photo_tables(
            photo_tables,
            node,
            doc=doc,
            leg_name=leg_name,
            project_path=project_path,
            remote_root=remote_root,
            pattern_dir=pattern_dir,
            matched=matched,
        )

    @staticmethod
    def _renumber_detail(table: Table, index: int) -> None:
        import re

        for row_i, row in enumerate(table.rows):
            cells = unique_cells(row)
            if not cells:
                continue
            raw = cell_text(cells[0])
            updated = re.sub(r"7\.\d+", f"7.{index}", raw, count=1)
            if updated != raw:
                force_black_text(cells[0], updated)

    def _fill_remark_cell(
        self,
        cell,
        node: TestNode,
        *,
        project_path: Optional[str],
        remote_root: Optional[str],
        leg_name: str,
    ) -> None:
        """Put note text, images, and tables into the template 备注 cell. Empty stays '/'."""
        parts = self._word._load_note_parts(node, project_path, leg_name, remote_root)
        if parts is None:
            set_blue_portion(cell, "/")
            return
        blocks, images, loaded = parts
        set_blue_portion(cell, blocks[0] if blocks else "")
        self._word.append_note_to_cell(cell, blocks, images, loaded)

    def _fill_equipment_rows(self, table: Table, node: TestNode) -> None:
        items = list(node.equipments or [])
        if not items and node.equipment_name:
            items = [TestEquipment(name=node.equipment_name)]
        eq_rows = []
        for idx, row in enumerate(table.rows):
            cells = unique_cells(row)
            if len(cells) >= 3 and "试验设备" in cell_text(cells[1]):
                eq_rows.append(idx)
        if len(eq_rows) < 2:
            return
        data_idx = eq_rows[1:]
        needed = len(items)
        while len(data_idx) < needed:
            src_i = data_idx[-1]
            insert_cloned_row_after(table, src_i)
            data_idx.append(src_i + 1)
        for row_i in reversed(data_idx[needed:]):
            delete_table_row(table, row_i)
        data_idx = data_idx[:needed]
        for offset, row_i in enumerate(data_idx):
            cells = unique_cells(table.rows[row_i])
            if len(cells) < 6:
                continue
            eq = items[offset]
            vals = [
                (eq.name or "").replace("\n", "") or "/",
                eq.model or "/",
                eq.code or "/",
                (eq.valid_date or "").strip() or "/",
            ]
            for col, val in enumerate(vals, start=2):
                self._write_blue(cells[col], val)

    @staticmethod
    def _photo_table_slot_count(table: Table) -> int:
        n = 0
        for cap_row in (2, 4, 6, 8):
            if cap_row >= len(table.rows):
                continue
            n += len(unique_cells(table.rows[cap_row]))
        return n

    def _ensure_photo_tables(
        self, tables: List[Table], needed: int, doc: Document
    ) -> List[Table]:
        if needed <= 0 or not tables:
            return list(tables)
        live = list(tables)
        proto = deepcopy(live[0]._tbl)
        while len(live) < needed:
            live.append(insert_cloned_table_after(live[-1], doc, proto))
        return live

    def _fill_photo_tables(
        self,
        tables: Sequence[Table],
        node: TestNode,
        *,
        doc: Document,
        leg_name: str,
        project_path: Optional[str],
        remote_root: Optional[str],
        pattern_dir: Optional[Path],
        matched: bool,
    ) -> None:
        photos = self._collect_photos(
            node, leg_name, project_path, remote_root, pattern_dir
        )
        live = list(tables)
        if photos and live:
            slots = self._photo_table_slot_count(live[0]) or 8
            needed = (len(photos) + slots - 1) // slots
            live = self._ensure_photo_tables(live, needed, doc)
        used = 0
        for table in live:
            title_cells = unique_cells(table.rows[0]) if table.rows else []
            if title_cells and not matched:
                name = (node.test_name or "").strip()
                force_black_text(title_cells[0], f"{name} 相关图片：")
            pairs: List[Tuple[int, int, List[Tuple]]] = []
            for cap_row in (2, 4, 6, 8):
                if cap_row >= len(table.rows):
                    continue
                img_row = cap_row - 1
                cap_cells = unique_cells(table.rows[cap_row])
                img_cells = unique_cells(table.rows[img_row])
                slots = []
                for col, cap_cell in enumerate(cap_cells):
                    img_cell = img_cells[col] if col < len(img_cells) else cap_cell
                    slots.append((cap_cell, img_cell))
                pairs.append((img_row, cap_row, slots))
            last_used_pair = -1
            last_pair_used = 0
            filled_here = 0
            for pair_i, (_img_row, _cap_row, slots) in enumerate(pairs):
                pair_used = 0
                for cap_cell, img_cell in slots:
                    remove_drawings(img_cell)
                    if used < len(photos):
                        path = photos[used]
                        set_photo_caption(
                            cap_cell, used + 1, caption_from_stem(path.stem)
                        )
                        self._put_picture(img_cell, path)
                        used += 1
                        filled_here += 1
                        pair_used += 1
                    else:
                        force_black_text(cap_cell, "")
                        force_black_text(img_cell, "")
                if pair_used:
                    last_used_pair = pair_i
                    last_pair_used = pair_used
            if filled_here == 0:
                delete_table(table)
                continue
            for pair_i in range(len(pairs) - 1, last_used_pair, -1):
                img_row, cap_row, _slots = pairs[pair_i]
                delete_table_row(table, cap_row)
                delete_table_row(table, img_row)
            if last_used_pair >= 0 and last_pair_used == 1:
                img_row, cap_row, slots = pairs[last_used_pair]
                if len(slots) > 1:
                    delete_row_cell(table.rows[img_row], -1)
                    delete_row_cell(table.rows[cap_row], -1)

    def _collect_photos(
        self,
        node: TestNode,
        leg_name: str,
        project_path: Optional[str],
        remote_root: Optional[str],
        pattern_dir: Optional[Path],
    ) -> List[Path]:
        from src.io.project_assets import iter_merged_export_photos

        photos: List[Path] = []
        if project_path and node.test_name and leg_name:
            for item in iter_merged_export_photos(
                Path(project_path),
                Path(remote_root) if remote_root else None,
                leg_name,
                node.test_name,
                order=getattr(node, "photo_album_order", None) or None,
                photo_file_order=getattr(node, "photo_file_order", None) or None,
                temps=self._export_temps,
            ):
                photos.append(item.path)
        photos.extend(self._pattern_photos(node.test_name or "", pattern_dir))
        return photos

    def _prefetch_export_photos(
        self,
        pairs: Sequence[Tuple],
        project_path: Optional[str],
        remote_root: Optional[str],
        pattern_dir: Optional[Path],
    ) -> None:
        from src.generators.embed_cache import prefetch_embed_paths
        from src.io.test_photos import list_sample_info_photos

        paths: List[Path] = []
        if project_path:
            try:
                paths.extend(list_sample_info_photos(Path(project_path)))
            except OSError:
                pass
        for leg, node in pairs:
            paths.extend(
                self._collect_photos(
                    node,
                    leg.leg_name,
                    project_path,
                    remote_root,
                    pattern_dir,
                )
            )
        prefetch_embed_paths(paths)

    @staticmethod
    def _pattern_photos(test_name: str, pattern_dir: Optional[Path]) -> List[Path]:
        if pattern_dir is None:
            pattern_dir = elp_data_pattern_directory()
        root = Path(pattern_dir)
        try:
            if not root.is_dir() or not test_name:
                return []
            folder = root / test_name
            if not folder.is_dir():
                folder = next(
                    (
                        child
                        for child in root.iterdir()
                        if child.is_dir() and elp_names_match(child.name, test_name)
                    ),
                    None,
                )
            if folder is None or not folder.is_dir():
                return []
            photos = [p for p in folder.iterdir() if is_image_file(p)]
        except OSError:
            return []
        photos.sort(key=lambda p: p.name.casefold())
        return photos

    def _put_picture(self, cell, path: Path) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        try:
            stream = self._word._prepare_embed_stream(path, PHOTO_WIDTH_IN)
            run = paragraph.add_run()
            run.add_picture(stream, width=Inches(PHOTO_WIDTH_IN))
        except Exception:
            paragraph.add_run(f"[无法插入图片: {path.name}]")

    @staticmethod
    def _write_blue(cell, value: str) -> None:
        if not set_blue_portion(cell, value):
            force_blue_text(cell, value)

    @staticmethod
    def _period(start: Optional[str], end: Optional[str]) -> str:
        a = WordGenerator._fmt_date(start)
        b = WordGenerator._fmt_date(end)
        if a and b:
            return f"{a}-{b}"
        return a or b or "/"

    @staticmethod
    def _conclusion(node: TestNode) -> str:
        return _node_conclusion_zh(node)

    @staticmethod
    def _result_text(node: TestNode) -> str:
        samples = [s for s in (node.samples or []) if s.sample_id]
        stds = node.result_table_standards()
        lines = []
        for sample in samples:
            desc = ""
            if stds:
                desc = (sample.desc_for(stds[0]) or stds[0].result_desc or "").strip()
            desc = desc or (sample.result_desc or node.result_desc or "").strip()
            lines.append(f"{sample.sample_id}：{desc or '/'}")
        return "\n".join(lines) if lines else "/"


def _run_is_blue(run) -> bool:
    from src.generators.elp_docx import run_is_blue

    return run_is_blue(run)


def _node_conclusion_zh(node: TestNode) -> str:
    engine = WordGenerator.__new__(WordGenerator)
    engine._report_language = "中文"
    return WordGenerator._node_conclusion(engine, node)

