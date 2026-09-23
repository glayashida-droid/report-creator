"""Fill original-record Word templates (原始记录) from a detail-dialog snapshot."""

from __future__ import annotations

import io
import re
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from src.generators.word_engine import (
    WordGenerator,
    _OR_SAMPLE_COL_DXA,
    _OR_TABLE_INDENT_DXA,
    _OR_TABLE_WIDTH_DXA,
)
from src.io.data_tables import (
    prepare_display_snapshot,
    read_preview_snapshot,
)
from src.io.network_sources import (
    local_fallback_root,
    original_data_sheet_directory,
)
from src.io.project_assets import (
    list_merged_attachment_refs,
    resolve_data_table_path,
)
from src.io.project_board import PROJECT_INTRANET_SHARE
from src.io.test_photos import TEST_GROUP_DIR, is_usable_test_name, test_dir
from src.models.project_state import (
    DataTableRef,
    TestEquipment,
    TestNode,
    TestStandard,
)

ORIGINAL_RECORD_TEMPLATE_NAME = "original_record_placeholders.docx"
RESULT_TABLE_MARKER = "{{RESULT_TABLE}}"
DATA_TABLES_MARKER = "{{DATA_TABLES}}"


@dataclass
class OriginalRecordData:
    application_no: str = ""
    test_item: str = ""
    test_method: str = ""
    sample_name: str = ""
    sample_ids: List[str] = field(default_factory=list)
    env_condition: str = ""
    start_date: str = ""
    end_date: str = ""
    share_path: str = ""
    equipments: List[TestEquipment] = field(default_factory=list)
    standards: List[TestStandard] = field(default_factory=list)
    tester_name: str = ""
    data_tables: List[DataTableRef] = field(default_factory=list)
    note_text: str = ""
    note_tables: List[DataTableRef] = field(default_factory=list)
    note_image_order: List[str] = field(default_factory=list)
    project_path: str = ""
    remote_root: str = ""
    leg_name: str = ""


def original_record_data_from_node(
    node: TestNode,
    *,
    leg_name: str,
    state=None,
    project_path: str = "",
    remote_root: str = "",
) -> OriginalRecordData:
    """Build original-record payload from a persisted test node (not a live dialog)."""
    app_no = ""
    sample_name = ""
    tester = ""
    share = ""
    if state is not None:
        fields = getattr(state, "application_fields", None) or {}
        app_no = (fields.get("申请单号") or getattr(state, "project_id", "") or "").strip()
        sample_name = (getattr(state, "sample_name", None) or "").strip()
        tester = (getattr(state, "tester_name", None) or "").strip()
        share = (getattr(state, "source_path", None) or "").strip()
        if not project_path:
            project_path = str(getattr(state, "project_path", "") or "")
        if not remote_root:
            remote_root = str(getattr(state, "source_path", "") or "")
    return OriginalRecordData(
        application_no=app_no,
        test_item=(node.test_name or "").strip(),
        test_method=node.resolved_test_method(),
        sample_name=sample_name,
        sample_ids=[
            s.sample_id.strip()
            for s in (node.samples or [])
            if (getattr(s, "sample_id", None) or "").strip()
        ],
        env_condition=node.resolved_env_condition(),
        start_date=(node.start_date or "").strip(),
        end_date=(node.end_date or "").strip(),
        share_path=share,
        equipments=list(node.equipments or []),
        standards=list(node.resolved_standards()),
        tester_name=tester,
        data_tables=list(node.data_tables or []),
        note_text=getattr(node, "note_text", "") or "",
        note_tables=list(getattr(node, "note_tables", None) or []),
        note_image_order=list(getattr(node, "note_image_order", None) or []),
        project_path=str(project_path or ""),
        remote_root=str(remote_root or ""),
        leg_name=(leg_name or "").strip(),
    )


def resolve_original_record_template(config=None) -> Path:
    """Local mirror first, then the configured original_data_sheet dir."""
    name = ORIGINAL_RECORD_TEMPLATE_NAME
    local = local_fallback_root() / "original_data_sheet" / name
    if local.is_file():
        return local
    try:
        folder = original_data_sheet_directory(config)
        candidate = folder / name
        if candidate.is_file():
            return candidate
    except Exception:
        pass
    raise FileNotFoundError(f"找不到原始记录模板: {name}")


def format_share_path_for_record(
    raw: str,
    *,
    smb_hint: str = PROJECT_INTRANET_SHARE,
) -> str:
    """Prefer a pasteable UNC path (\\\\server\\share\\...) for Explorer / cross-OS notes."""
    text = (raw or "").strip().rstrip("/\\")
    if not text:
        return ""
    if text.startswith("smb://"):
        return _smb_url_to_unc(text)
    if text.startswith("\\\\") or text.startswith("//"):
        parts = [p for p in text.replace("/", "\\").split("\\") if p]
        return "\\\\" + "\\".join(parts) if parts else text
    if text.startswith("/Volumes/"):
        host = _smb_host(smb_hint)
        parts = [p for p in text.split("/") if p]
        # parts[0] == Volumes, parts[1] == share, …
        if host and len(parts) >= 2:
            return "\\\\" + host + "\\" + "\\".join(parts[1:])
    return text


def _smb_host(smb_hint: str) -> str:
    text = (smb_hint or "").strip()
    if text.startswith("smb://"):
        rest = text[6:]
        return rest.split("/", 1)[0] if rest else ""
    if text.startswith("\\\\"):
        parts = [p for p in text.replace("/", "\\").split("\\") if p]
        return parts[0] if parts else ""
    return ""


def _smb_url_to_unc(text: str) -> str:
    without = text[6:]
    parts = [p for p in without.split("/") if p]
    if not parts:
        return text
    return "\\\\" + parts[0] + ("\\" + "\\".join(parts[1:]) if len(parts) > 1 else "")


def format_id_range(ids: Sequence[str]) -> str:
    clean = [i.strip() for i in ids if i and str(i).strip()]
    if not clean:
        return ""
    if len(clean) == 1:
        return clean[0]

    def split_tail(s: str):
        m = re.match(r"^(.*?)(\d+)$", s)
        return (m.group(1), m.group(2)) if m else (s, "")

    prefix0, num0 = split_tail(clean[0])
    prefix1, num1 = split_tail(clean[-1])
    if prefix0 and prefix0 == prefix1 and num0 and num1 and len(clean) > 1:
        return f"{clean[0]}~{clean[-1]}"
    return "、".join(clean)


def format_test_period(start: Optional[str], end: Optional[str]) -> str:
    a = _fmt_date(start)
    b = _fmt_date(end)
    if a and b:
        return f"{a}~{b}"
    return a or b or ""


def _fmt_date(value: Optional[str]) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text[:10], fmt).strftime("%Y.%m.%d")
        except ValueError:
            continue
    return text.replace("-", ".").replace("/", ".")


def stack_standard_blocks(
    standards: Sequence[TestStandard],
    attr: str,
) -> str:
    """Stack standard field texts; separate with 试验名称 when multiple."""
    items = list(standards or [])
    if not items:
        return ""
    blocks: List[str] = []
    for std in items:
        title = (std.test_name or std.condition_title() or "").strip()
        body = (getattr(std, attr, None) or "").strip()
        if len(items) == 1:
            if body:
                blocks.append(body)
            continue
        if body:
            blocks.append(f"{title}\n{body}" if title else body)
        elif title:
            blocks.append(title)
    return "\n\n".join(blocks)


def build_original_record_document(
    data: OriginalRecordData,
    *,
    template_path: Path | str | None = None,
) -> Document:
    """Fill one original-record document in memory."""
    template = Path(template_path) if template_path else resolve_original_record_template()
    doc = Document(str(template))
    _fill_original_record(doc, data, template)
    return doc


def generate_original_record(
    data: OriginalRecordData,
    output_path: Path | str,
    *,
    template_path: Path | str | None = None,
) -> Path:
    template = Path(template_path) if template_path else resolve_original_record_template()
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = build_original_record_document(data, template_path=template)
    doc.save(str(out))
    return out


def _fill_original_record(doc: Document, data: OriginalRecordData, template: Path) -> None:
    _expand_equipment_rows(doc, data.equipments or [])
    _expand_result_tables(doc, data.standards or [], data.sample_ids or [])
    _remove_marker_paragraph(doc, RESULT_TABLE_MARKER)
    _fill_data_tables_section(doc, data, template_path=template)
    _fill_test_parameters_section(doc, data.standards or [], template_path=template)
    _fill_note_section(doc, data, template_path=template)

    placeholders = {
        "{{申请单编号}}": data.application_no or "",
        "{{检验项目}}": data.test_item or "",
        "{{检测依据}}": data.test_method or "",
        "{{样品名称}}": data.sample_name or "",
        "{{样品编号}}": format_id_range(data.sample_ids),
        "{{试验环境}}": data.env_condition or "",
        "{{测试数量}}": str(len([s for s in (data.sample_ids or []) if str(s).strip()])),
        "{{检验时间}}": format_test_period(data.start_date, data.end_date),
        "{{公盘地址}}": format_share_path_for_record(data.share_path),
        "{{评价要求}}": stack_standard_blocks(data.standards, "evaluation_req"),
        "{{检测人}}": data.tester_name or "",
    }
    _replace_everywhere(doc, placeholders)
    # leftover structural placeholders if no equipment / no standards
    _replace_everywhere(
        doc,
        {
            "{{测试参数}}": "",
            "{{设备名称}}": "",
            "{{设备型号}}": "",
            "{{设备编号}}": "",
            "{{检定有效期}}": "",
            "{{试验项目}}": "",
            "{{结果样品编号}}": "",
            "{{DATA_TABLES}}": "/",
        },
    )


def _existing_dir(raw: Optional[Path | str]) -> Optional[Path]:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    path = Path(text)
    return path if path.is_dir() else None


def resolve_original_record_folder(
    project_dir: Optional[Path | str],
    remote_root: Optional[Path | str] = None,
    *,
    leg_name: str = "",
    test_item: str = "",
) -> Optional[Path]:
    """试验目录（与照片夹同级的父目录）。公盘可达则用公盘，否则本地镜像。"""
    root = _existing_dir(remote_root) or _existing_dir(project_dir)
    leg = (leg_name or "").strip()
    test = (test_item or "").strip()
    if root is None or not leg or not is_usable_test_name(test):
        return None
    return test_dir(root, leg, test)


def default_output_path(
    folder: Optional[Path],
    *,
    application_no: str = "",
    test_item: str = "",
) -> Path:
    stem_bits = [p for p in ((application_no or "").strip(), _safe_stem(test_item)) if p]
    stem = "_".join(stem_bits) if stem_bits else "original_record"
    stem = f"{stem}_原始记录"
    dest = folder if folder is not None else Path(".scratch")
    dest.mkdir(parents=True, exist_ok=True)
    out = dest / f"{stem}.docx"
    if not out.exists():
        return out
    n = 2
    while True:
        candidate = dest / f"{stem}-{n}.docx"
        if not candidate.exists():
            return candidate
        n += 1


def export_node_original_record(
    data: OriginalRecordData,
    *,
    template_path: Path | str,
) -> Path:
    """Write one original-record Word file beside the test albums."""
    if not is_usable_test_name(data.test_item):
        raise ValueError("请先选择试验名称")
    if not (data.leg_name or "").strip():
        raise ValueError("缺少 Leg 名称")
    folder = resolve_original_record_folder(
        data.project_path or None,
        data.remote_root or None,
        leg_name=data.leg_name,
        test_item=data.test_item,
    )
    if folder is None:
        raise FileNotFoundError("找不到项目目录（公盘与本地均不可用）")
    out_path = default_output_path(
        folder,
        application_no=data.application_no,
        test_item=data.test_item,
    )
    return generate_original_record(data, out_path, template_path=template_path)


def resolve_combined_original_record_folder(
    project_dir: Optional[Path | str],
    remote_root: Optional[Path | str] = None,
) -> Optional[Path]:
    """3.测试组/，与各试验目录同级。公盘可达则用公盘，否则本地镜像。"""
    root = _existing_dir(remote_root) or _existing_dir(project_dir)
    if root is None:
        return None
    return root / TEST_GROUP_DIR


def export_leg_original_record(
    records: Sequence[OriginalRecordData],
    *,
    template_path: Path | str,
) -> Path:
    """Write one Word file that stacks every test on this leg, with a page break between them."""
    items = list(records or [])
    if not items:
        raise ValueError("没有可打印的试验")
    for data in items:
        if not is_usable_test_name(data.test_item):
            raise ValueError("请先选择试验名称")
    leg = (items[0].leg_name or "").strip()
    if not leg:
        raise ValueError("缺少 Leg 名称")
    folder = resolve_combined_original_record_folder(
        items[0].project_path or None,
        items[0].remote_root or None,
    )
    if folder is None:
        raise FileNotFoundError("找不到项目目录（公盘与本地均不可用）")
    out_path = default_output_path(
        folder,
        application_no=items[0].application_no,
        test_item=leg,
    )
    base = build_original_record_document(items[0], template_path=template_path)
    for data in items[1:]:
        extra = build_original_record_document(data, template_path=template_path)
        append_original_record(base, extra)
    base.save(str(out_path))
    return out_path


def append_original_record(base: Document, extra: Document) -> None:
    """Append another filled original record after a page break. Images keep their bytes."""
    body = base.element.body
    sect = body.find(qn("w:sectPr"))
    _insert_body_child(body, sect, _page_break_paragraph())
    cache: Dict[str, str] = {}
    for child in list(extra.element.body):
        if child.tag == qn("w:sectPr"):
            continue
        cloned = deepcopy(child)
        _remap_relationships(cloned, extra.part, base.part, cache)
        _insert_body_child(body, sect, cloned)


def _insert_body_child(body, sect, element) -> None:
    if sect is not None:
        sect.addprevious(element)
    else:
        body.append(element)


def _page_break_paragraph():
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    return paragraph


def _remap_relationships(element, src_part, dest_part, cache: Dict[str, str]) -> None:
    attrs = (qn("r:embed"), qn("r:link"), qn("r:id"))
    for node in element.iter():
        for attr in attrs:
            old = node.get(attr)
            if not old or old not in src_part.rels:
                continue
            if old not in cache:
                cache[old] = _rehome_relationship(src_part.rels[old], dest_part)
            node.set(attr, cache[old])


def _rehome_relationship(rel, dest_part) -> str:
    if rel.is_external:
        return dest_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
    if rel.reltype == RT.IMAGE:
        new_rid, _image = dest_part.get_or_add_image(io.BytesIO(rel.target_part.blob))
        return new_rid
    return dest_part.relate_to(rel.target_part, rel.reltype)


def _safe_stem(text: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", (text or "").strip())
    return cleaned.strip("._")[:80]


def _fill_note_section(doc: Document, data: OriginalRecordData, *, template_path: Path) -> None:
    """Replace the template underline under 备注(note). Leave it when the note is empty."""
    anchor = _find_note_underline(doc)
    if anchor is None:
        return
    node = TestNode(
        test_name=(data.test_item or "").strip() or "请选择试验...",
        note_text=data.note_text or "",
        note_tables=list(data.note_tables or []),
        note_image_order=list(data.note_image_order or []),
    )
    engine = WordGenerator(str(template_path))
    inserted = engine._insert_note_section(
        doc,
        anchor,
        node,
        data.project_path or None,
        data.leg_name or "",
        remote_root=data.remote_root or None,
        label=None,
    )
    if inserted:
        engine._delete_paragraph(anchor)


def _find_note_underline(doc: Document) -> Optional[Paragraph]:
    paras = list(doc.paragraphs)
    for index, paragraph in enumerate(paras):
        text = (paragraph.text or "").strip()
        folded = text.casefold()
        if not text.startswith("备注") or "note" not in folded:
            continue
        for follower in paras[index + 1 : index + 4]:
            line = (follower.text or "").strip()
            if line and set(line) <= {"_"}:
                return follower
    return None


def _fill_test_parameters_section(
    doc: Document,
    standards: Sequence[TestStandard],
    *,
    template_path: Path,
) -> None:
    """Fill {{测试参数}} like report 检测条件: text, then blank, then library images."""
    anchor = _find_paragraph(doc, "{{测试参数}}")
    if anchor is None:
        return

    stds = list(standards or [])
    if not stds:
        _replace_text_in_paragraph(anchor, {"{{测试参数}}": ""})
        return

    engine = WordGenerator(str(template_path))
    node = TestNode(test_name="")
    node.apply_standards(stds)
    engine._insert_condition_section(doc, anchor, node)
    engine._delete_paragraph(anchor)


def _fill_data_tables_section(
    doc: Document,
    data: OriginalRecordData,
    *,
    template_path: Path,
) -> None:
    """Replace {{DATA_TABLES}} with stacked tables (titles) or '/' when empty."""
    anchor = _find_paragraph(doc, DATA_TABLES_MARKER)
    if anchor is None:
        return

    local = Path(data.project_path) if (data.project_path or "").strip() else None
    remote = Path(data.remote_root) if (data.remote_root or "").strip() else None
    leg = (data.leg_name or "").strip()
    test_name = (data.test_item or "").strip()

    refs: List[DataTableRef] = []
    if (local is not None or remote is not None) and leg and test_name:
        refs = list_merged_attachment_refs(local, remote, leg, test_name)
    if not refs:
        refs = list(data.data_tables or [])

    # Preserve per-ref limit checkbox from the live dialog list.
    limit_flags = {
        r.relative_path: bool(getattr(r, "include_limit_row_in_report", False))
        for r in (data.data_tables or [])
    }
    for r in refs:
        if r.relative_path in limit_flags:
            r.include_limit_row_in_report = limit_flags[r.relative_path]

    loaded = []
    for ref in refs:
        try:
            path = resolve_data_table_path(local, remote, ref.relative_path)
            if path is None:
                continue
            snap = read_preview_snapshot(path)
        except Exception:
            continue
        if not snap.values:
            continue
        include_limit = bool(getattr(ref, "include_limit_row_in_report", False))
        snap = prepare_display_snapshot(snap, include_limit_row=include_limit)
        if not snap.values:
            continue
        loaded.append((ref, path, snap))

    if not loaded:
        _replace_text_in_paragraph(anchor, {DATA_TABLES_MARKER: "/"})
        return

    engine = WordGenerator(str(template_path))
    for ref, path, snap in loaded:
        title = Path(ref.title or path.stem).stem.strip()
        engine._render_data_table_snapshot(
            doc,
            anchor,
            snap,
            title=title,
            sample_dxa=_OR_SAMPLE_COL_DXA,
            total_dxa=_OR_TABLE_WIDTH_DXA,
            indent_dxa=_OR_TABLE_INDENT_DXA,
        )

    engine._delete_paragraph(anchor)


def _find_paragraph(doc: Document, marker: str) -> Optional[Paragraph]:
    for p in doc.paragraphs:
        if marker in (p.text or ""):
            return p
    for table in doc.tables:
        hit = _find_paragraph_in_table(table, marker)
        if hit is not None:
            return hit
    return None


def _find_paragraph_in_table(table: Table, marker: str) -> Optional[Paragraph]:
    for row in table.rows:
        for tc in row._tr.tc_lst:
            cell = _Cell(tc, table)
            for p in cell.paragraphs:
                if marker in (p.text or ""):
                    return p
            for nested in cell.tables:
                hit = _find_paragraph_in_table(nested, marker)
                if hit is not None:
                    return hit
    return None


# --------------------------------------------------------------------------- XML helpers


def _set_vmerge(cell: _Cell, value: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("w:vMerge")):
        tcPr.remove(el)
    vm = OxmlElement("w:vMerge")
    if value == "restart":
        vm.set(qn("w:val"), "restart")
    tcPr.append(vm)


def _clear_vmerge_tc(tc) -> None:
    tcPr = tc.tcPr
    if tcPr is None:
        return
    for el in tcPr.findall(qn("w:vMerge")):
        tcPr.remove(el)


def _set_cell_text(cell: _Cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.text = text
        return
    first = paras[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    if text:
        first.add_run(text)
    for p in paras[1:]:
        parent = p._element.getparent()
        if parent is not None:
            parent.remove(p._element)


def _replace_everywhere(doc: Document, placeholders: dict) -> None:
    if not placeholders:
        return
    for p in doc.paragraphs:
        _replace_text_in_paragraph(p, placeholders)
    for table in doc.tables:
        _replace_in_table(table, placeholders)


def _replace_in_table(table: Table, placeholders: dict) -> None:
    for row in table.rows:
        for tc in row._tr.tc_lst:
            cell = _Cell(tc, table)
            for p in cell.paragraphs:
                _replace_text_in_paragraph(p, placeholders)
            for nested in cell.tables:
                _replace_in_table(nested, placeholders)


def _replace_text_in_paragraph(paragraph: Paragraph, placeholders: dict) -> None:
    text = paragraph.text
    if not text:
        return
    new_text = text
    for key, val in placeholders.items():
        if key in new_text:
            new_text = new_text.replace(key, str(val))
    if new_text == text:
        return
    # Prefer per-run when a placeholder sits in one run
    changed = False
    for run in paragraph.runs:
        rt = run.text or ""
        nt = rt
        for key, val in placeholders.items():
            if key in nt:
                nt = nt.replace(key, str(val))
        if nt != rt:
            run.text = nt
            changed = True
    if changed and paragraph.text == new_text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def _remove_marker_paragraph(doc: Document, marker: str) -> None:
    for p in list(doc.paragraphs):
        if marker in (p.text or ""):
            el = p._element
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
            return


# --------------------------------------------------------------------------- equipment


def _expand_equipment_rows(doc: Document, equipments: Sequence[TestEquipment]) -> None:
    table = _find_equipment_table(doc)
    if table is None:
        return
    data_idx = _find_row_with_placeholder(table, "{{设备名称}}")
    if data_idx is None:
        return
    items = list(equipments or [])
    if not items:
        return

    template_tr = table.rows[data_idx]._tr
    rows_trs = [template_tr]
    for _ in range(1, len(items)):
        new_tr = deepcopy(template_tr)
        for tc in new_tr.tc_lst:
            _clear_vmerge_tc(tc)
        rows_trs[-1].addnext(new_tr)
        rows_trs.append(new_tr)

    data_idx = _find_row_with_placeholder(table, "{{设备名称}}")
    assert data_idx is not None
    label_row_idx = data_idx - 1 if data_idx > 0 else None

    for i, eq in enumerate(items):
        row = table.rows[data_idx + i]
        tcs = row._tr.tc_lst
        # Template layout: label | name | model(span) | code(span) | valid
        if len(tcs) >= 5:
            name_c = _Cell(tcs[1], table)
            model_c = _Cell(tcs[2], table)
            code_c = _Cell(tcs[3], table)
            valid_c = _Cell(tcs[4], table)
            left = _Cell(tcs[0], table)
        else:
            name_c = _Cell(tcs[1], table)
            model_c = _Cell(tcs[2], table)
            code_c = _Cell(tcs[min(3, len(tcs) - 2)], table)
            valid_c = _Cell(tcs[-1], table)
            left = _Cell(tcs[0], table)
        _set_cell_text(name_c, eq.name or "")
        _set_cell_text(model_c, eq.model or "")
        _set_cell_text(code_c, eq.code or "")
        _set_cell_text(valid_c, eq.valid_date or "")
        _set_cell_text(left, "")
        _set_vmerge(left, "continue")

    if label_row_idx is not None:
        left_hdr = _Cell(table.rows[label_row_idx]._tr.tc_lst[0], table)
        _set_vmerge(left_hdr, "restart")


def _find_equipment_table(doc: Document) -> Optional[Table]:
    for table in doc.tables:
        for row in table.rows:
            for tc in row._tr.tc_lst:
                text = _Cell(tc, table).text or ""
                if "使用设备" in text or "{{设备名称}}" in text:
                    return table
    return None


def _find_row_with_placeholder(table: Table, needle: str) -> Optional[int]:
    for i, row in enumerate(table.rows):
        for tc in row._tr.tc_lst:
            if needle in (_Cell(tc, table).text or ""):
                return i
    return None


# --------------------------------------------------------------------------- result tables


def _expand_result_tables(
    doc: Document,
    standards: Sequence[TestStandard],
    sample_ids: Sequence[str],
) -> None:
    proto = _find_result_prototype_table(doc)
    if proto is None:
        return
    ids = [s.strip() for s in sample_ids if s and str(s).strip()]
    stds = list(standards or [])
    titles = [
        (s.test_name or s.condition_title() or "").strip() for s in stds
    ] or [""]

    result_tables = [proto]
    for _ in range(1, len(titles)):
        src = result_tables[-1]
        new_tbl = deepcopy(proto._tbl)
        src._tbl.addnext(new_tbl)
        result_tables.append(Table(new_tbl, proto._parent))

    for title, table in zip(titles, result_tables):
        _fill_one_result_table(table, title, ids)


def _table_has_any(table: Table, needles: Iterable[str]) -> bool:
    parts = []
    for row in table.rows:
        for tc in row._tr.tc_lst:
            parts.append(_Cell(tc, table).text)
    blob = "\n".join(parts)
    return any(n in blob for n in needles)


def _find_result_prototype_table(doc: Document) -> Optional[Table]:
    for table in doc.tables:
        if _table_has_any(table, ("{{试验项目}}", "{{结果样品编号}}")):
            return table
    if len(doc.tables) >= 2:
        return doc.tables[1]
    return None


def _vmerge_column_indexes(*rows) -> List[int]:
    """Column indexes that already use w:vMerge in any of the given rows."""
    cols: set[int] = set()
    for row in rows:
        for ci, tc in enumerate(row._tr.tc_lst):
            tcPr = tc.tcPr
            if tcPr is not None and tcPr.find(qn("w:vMerge")) is not None:
                cols.add(ci)
    return sorted(cols)


def _apply_sample_block_vmerge(
    table: Table, start: int, block_len: int, merge_cols: Sequence[int]
) -> None:
    """Restart/continue vertical merges inside one 试验前/中/后 sample block."""
    for offset in range(block_len):
        tcs = table.rows[start + offset]._tr.tc_lst
        val = "restart" if offset == 0 else "continue"
        for ci in merge_cols:
            if ci < len(tcs):
                _set_vmerge(_Cell(tcs[ci], table), val)


def _fill_one_result_table(table: Table, test_name: str, sample_ids: Sequence[str]) -> None:
    _replace_in_table(table, {"{{试验项目}}": test_name or ""})
    block_starts = _sample_block_starts(table)
    if not block_starts:
        return
    ids = list(sample_ids) or [""]
    first_start = block_starts[0]
    block_len = 3
    proto_rows = [table.rows[first_start + o] for o in range(block_len)]
    merge_cols = _vmerge_column_indexes(*proto_rows)
    if 0 not in merge_cols:
        merge_cols = [0, *merge_cols]

    # Clone sample blocks first (while the prototype still has the sample placeholder).
    while len(_sample_block_starts(table)) < len(ids):
        starts = _sample_block_starts(table)
        last_start = starts[-1]
        anchor = table.rows[last_start + block_len - 1]._tr
        for offset in range(block_len):
            src = table.rows[first_start + offset]._tr
            new_tr = deepcopy(src)
            # Drop copied vMerge so clones don't continue the previous sample;
            # re-apply per block after fill, using the prototype's merged columns.
            for tc in new_tr.tc_lst:
                _clear_vmerge_tc(tc)
            anchor.addnext(new_tr)
            anchor = new_tr

    starts = _sample_block_starts(table)
    for i, sid in enumerate(ids):
        if i >= len(starts):
            break
        start = starts[i]
        for offset in range(block_len):
            cell = _Cell(table.rows[start + offset]._tr.tc_lst[0], table)
            _set_cell_text(cell, sid if offset == 0 else "")
        _apply_sample_block_vmerge(table, start, block_len, merge_cols)


def _sample_block_starts(table: Table) -> List[int]:
    """Row indices that begin a 试验前/中/后 sample block."""
    starts: List[int] = []
    for i, row in enumerate(table.rows):
        texts = []
        for tc in row._tr.tc_lst:
            texts.append(_Cell(tc, table).text.strip())
        if any(t == "试验前" or t.startswith("试验前") for t in texts):
            starts.append(i)
    return starts
