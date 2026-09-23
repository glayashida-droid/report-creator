"""Low-level python-docx helpers for the Geely ELP template.

Blue runs in the template mark fill slots. Generated reports write those
slots in black; leftover template blue is painted black at the end.
"""

from __future__ import annotations

from copy import deepcopy
from typing import Iterable, List, NamedTuple, Optional, Sequence

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

BLUE_VALS = {"0000FF", "0070C0", "00B0F0"}
BLACK_VAL = "000000"
CHECK_MARKS = {"√", "✓", "☑"}
NSMAP_R = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
NSMAP_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def unique_cells(row) -> list:
    seen = set()
    out = []
    for cell in row.cells:
        cid = id(cell._tc)
        if cid in seen:
            continue
        seen.add(cid)
        out.append(cell)
    return out


def cell_text(cell) -> str:
    return "".join((p.text or "") for p in cell.paragraphs).strip()


def run_color_val(run) -> str:
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        return ""
    color = rPr.find(qn("w:color"))
    if color is None:
        return ""
    return (color.get(qn("w:val")) or "").upper()


def run_is_blue(run) -> bool:
    return run_color_val(run) in BLUE_VALS


def run_has_drawing(run) -> bool:
    """True if the run carries a picture/shape. Assigning run.text would delete it."""
    for child in run._r:
        tag = child.tag.split("}")[-1]
        if tag in {"drawing", "pict", "AlternateContent", "object"}:
            return True
    return False


def cell_has_blue(cell) -> bool:
    for p in cell.paragraphs:
        for run in p.runs:
            if run_is_blue(run) and (run.text or "").strip():
                return True
    return False


def set_run_color(run, val: str) -> None:
    rPr = run._r.get_or_add_rPr()
    color = rPr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rPr.append(color)
    color.set(qn("w:val"), val)


def _paint_rpr_blue_to_black(rPr) -> None:
    color = rPr.find(qn("w:color"))
    if color is None:
        return
    val = (color.get(qn("w:val")) or "").upper()
    if val in BLUE_VALS:
        color.set(qn("w:val"), BLACK_VAL)


def _story_elements(doc) -> List:
    elements = [doc.element]
    seen = {id(doc.element)}
    try:
        rels = doc.part.rels.values()
    except Exception:
        return elements
    for rel in rels:
        reltype = getattr(rel, "reltype", "") or ""
        if "header" not in reltype and "footer" not in reltype:
            continue
        part = getattr(rel, "target_part", None)
        element = getattr(part, "element", None) if part is not None else None
        if element is None or id(element) in seen:
            continue
        seen.add(id(element))
        elements.append(element)
    return elements


def paint_blue_runs_black(doc) -> None:
    """Recolor leftover template-blue slots in the generated document."""
    for root in _story_elements(doc):
        for rPr in root.iter(qn("w:rPr")):
            _paint_rpr_blue_to_black(rPr)


def write_runs(paragraph, pieces: Sequence[tuple[str, Optional[str]]]) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    for text, color in pieces:
        run = paragraph.add_run(text)
        if color:
            set_run_color(run, color)


def clear_first_line_indent(paragraph) -> None:
    """Drop 首行缩进 so filled lines start at the left edge of the cell."""
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return
    for attr in (qn("w:firstLine"), qn("w:firstLineChars")):
        if attr in ind.attrib:
            del ind.attrib[attr]
    if not ind.attrib:
        pPr.remove(ind)


def fill_result_slot(cell, text: str) -> None:
    """Write 试验结果. Drop the template's exact row height so the row follows the text."""
    if not set_blue_portion(cell, text):
        force_black_text(cell, text)
    for paragraph in list(cell.paragraphs):
        clear_first_line_indent(paragraph)
    _drop_empty_paragraphs(cell)
    _unlock_row_height(cell)


def _unlock_row_height(cell) -> None:
    tr = cell._tc.getparent()
    if tr is None:
        return
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        return
    for height in list(trPr.findall(qn("w:trHeight"))):
        trPr.remove(height)


def _drop_empty_paragraphs(cell) -> None:
    """Remove template leftover paragraphs (the 2# / 3# stubs) after the text is written."""
    for paragraph in list(cell.paragraphs):
        if len(cell.paragraphs) <= 1:
            return
        if (paragraph.text or "").strip():
            continue
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def set_blue_portion(cell, text: str) -> bool:
    """Replace template-blue runs with filled text (black). Leave all-black cells."""
    blue_runs = []
    black_text = False
    for p in cell.paragraphs:
        for run in p.runs:
            if run_is_blue(run):
                blue_runs.append(run)
            elif (run.text or "").strip():
                black_text = True
    value = "" if text is None else str(text)
    if blue_runs:
        blue_runs[0].text = value
        set_run_color(blue_runs[0], BLACK_VAL)
        for run in blue_runs[1:]:
            run.text = ""
            set_run_color(run, BLACK_VAL)
        return True
    if black_text:
        return False
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    if paragraph is None:
        return False
    write_runs(paragraph, [(value, BLACK_VAL)])
    for extra in list(cell.paragraphs)[1:]:
        extra.text = ""
    return True


def force_blue_text(cell, text: str) -> None:
    """Overwrite a cell with filled text (black in the generated report)."""
    force_black_text(cell, text)


def force_black_text(cell, text: str) -> None:
    value = "" if text is None else str(text)
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    if paragraph is None:
        return
    write_runs(paragraph, [(value, None)])
    for extra in list(cell.paragraphs)[1:]:
        extra.text = ""


def set_photo_caption(cell, index: int, description: str) -> None:
    label = f"图片{index}："
    desc = (description or "").strip() or "/"
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    if paragraph is None:
        return
    write_runs(paragraph, [(label, None), (desc, None)])
    for extra in list(cell.paragraphs)[1:]:
        extra.text = ""


def delete_table_row(table: Table, index: int) -> None:
    table._tbl.remove(table.rows[index]._tr)


def delete_table(table: Table) -> None:
    tbl = table._tbl
    parent = tbl.getparent()
    if parent is not None:
        parent.remove(tbl)


def insert_cloned_table_after(table: Table, doc, prototype_tbl) -> Table:
    """Insert a deepcopy of prototype_tbl after ``table``, with a spacer paragraph."""
    new_tbl = deepcopy(prototype_tbl)
    spacer = OxmlElement("w:p")
    table._tbl.addnext(new_tbl)
    new_tbl.addprevious(spacer)
    return wrap_table(new_tbl, doc)


def delete_row_cell(row, index: int) -> None:
    """Drop one tc from a row. Remaining cells keep their own width (no gridSpan)."""
    tcs = list(row._tr.findall(qn("w:tc")))
    if index < 0:
        index += len(tcs)
    if len(tcs) < 2 or not (0 <= index < len(tcs)):
        return
    row._tr.remove(tcs[index])
    for tc in row._tr.findall(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            continue
        span = tcPr.find(qn("w:gridSpan"))
        if span is not None:
            tcPr.remove(span)


def clone_table_row(table: Table, source_index: int) -> None:
    src = table.rows[source_index]._tr
    table._tbl.append(deepcopy(src))


def insert_cloned_row_after(table: Table, source_index: int) -> None:
    src = table.rows[source_index]._tr
    src.addnext(deepcopy(src))


def remove_drawings(cell) -> None:
    for drawing in list(cell._tc.iter(qn("w:drawing"))):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)
    for pict in list(cell._tc.iter(qn("w:pict"))):
        parent = pict.getparent()
        if parent is not None:
            parent.remove(pict)


def paragraph_style_name(paragraph: Paragraph) -> str:
    try:
        style = paragraph.style
        return (style.name if style is not None else "") or ""
    except Exception:
        return ""


def iter_body_children(doc) -> List:
    return [child for child in doc.element.body if child.tag != qn("w:sectPr")]


def wrap_paragraph(element, doc) -> Paragraph:
    return Paragraph(element, doc)


def wrap_table(element, doc) -> Table:
    return Table(element, doc)


class GridCell(NamedTuple):
    text: str
    span: int
    vmerge: Optional[str]


def pad_grid(rows: Sequence[Sequence[str]]) -> List[List[str]]:
    width = max((len(row) for row in rows), default=0)
    return [list(row) + [""] * (width - len(row)) for row in rows]


class HGroup(NamedTuple):
    text: str
    span: int
    merge_up: bool


def _is_tick_row(row: Sequence[str]) -> bool:
    nonempty = [(cell or "").strip() for cell in row if (cell or "").strip()]
    return bool(nonempty) and all(cell in CHECK_MARKS for cell in nonempty)


def _h_groups(row: Sequence[str], *, attach_empties: bool) -> List[HGroup]:
    if not attach_empties:
        return [HGroup(cell or "", 1, False) for cell in row]
    groups: List[HGroup] = []
    i = 0
    n = len(row)
    while i < n:
        text = row[i] or ""
        span = 1
        i += 1
        if text.strip():
            while i < n and not (row[i] or "").strip():
                span += 1
                i += 1
        else:
            while i < n and not (row[i] or "").strip():
                span += 1
                i += 1
        groups.append(HGroup(text, span, False))
    return groups


def _map_onto_groups(row: Sequence[str], groups: Sequence[HGroup]) -> List[HGroup]:
    col = 0
    out: List[HGroup] = []
    for index, prev in enumerate(groups):
        chunk = list(row[col : col + prev.span])
        value = next((cell for cell in chunk if (cell or "").strip()), "")
        merge_up = index == 0 and not (value or "").strip()
        out.append(HGroup(value, prev.span, merge_up))
        col += prev.span
    return out


def grid_row_groups(rows: Sequence[Sequence[str]]) -> List[List[HGroup]]:
    """Infer horizontal merges from a PDF-extracted grid."""
    padded = pad_grid(rows)
    all_groups: List[List[HGroup]] = []
    prev: List[HGroup] = []
    for row in padded:
        if _is_tick_row(row) and prev:
            groups = _map_onto_groups(row, prev)
        elif any((cell or "").strip() in CHECK_MARKS for cell in row):
            groups = [HGroup(cell or "", 1, False) for cell in row]
        else:
            groups = []
            col = 0
            if prev:
                for item in prev:
                    chunk = row[col : col + item.span]
                    if chunk and all(not (cell or "").strip() for cell in chunk):
                        groups.append(HGroup("", item.span, True))
                        col += item.span
                    else:
                        break
            groups.extend(_h_groups(row[col:], attach_empties=True))
        all_groups.append(groups)
        prev = groups
    return all_groups


def _grid_col_widths(table: Table, ncols: int) -> List[int]:
    grid = table._tbl.find(qn("w:tblGrid"))
    widths: List[int] = []
    if grid is not None:
        for col in grid:
            try:
                widths.append(int(col.get(qn("w:w")) or 0))
            except (TypeError, ValueError):
                widths.append(0)
    if len(widths) == ncols and all(w > 0 for w in widths):
        return widths
    total = sum(widths) if widths else 9000
    each = max(int(total / max(ncols, 1)), 200)
    return [each] * ncols


def _set_tbl_grid(table: Table, widths: Sequence[int]) -> None:
    tbl = table._tbl
    old = tbl.find(qn("w:tblGrid"))
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width)))
        grid.append(col)
    if old is not None:
        old.getparent().replace(old, grid)
    else:
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            tblPr.addnext(grid)
        else:
            tbl.insert(0, grid)


def _make_grid_tc(
    text: str,
    span: int,
    width: int,
    vmerge: Optional[str],
) -> OxmlElement:
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(int(width)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    if span > 1:
        grid_span = OxmlElement("w:gridSpan")
        grid_span.set(qn("w:val"), str(span))
        tcPr.append(grid_span)
    if vmerge:
        vm = OxmlElement("w:vMerge")
        if vmerge != "continue":
            vm.set(qn("w:val"), vmerge)
        tcPr.append(vm)
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    tcPr.append(v_align)
    tc.append(tcPr)

    paragraph = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    if (text or "").strip() in CHECK_MARKS:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        pPr.append(jc)
    paragraph.append(pPr)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(_XML_SPACE, "preserve")
    t.text = text or ""
    run.append(t)
    paragraph.append(run)
    tc.append(paragraph)
    return tc


def replace_table_from_grid(table: Table, rows: Sequence[Sequence[str]]) -> None:
    """Replace a Word table's rows with a PDF-extracted grid, keeping borders."""
    if not rows:
        return
    padded = pad_grid(rows)
    ncols = len(padded[0])
    grouped = grid_row_groups(padded)
    specs: List[List[GridCell]] = []
    for r_i, groups in enumerate(grouped):
        row_specs: List[GridCell] = []
        col = 0
        for group in groups:
            vmerge = None
            if group.merge_up:
                vmerge = "continue"
            elif r_i + 1 < len(grouped):
                n_col = 0
                for nxt in grouped[r_i + 1]:
                    if n_col == col and nxt.span == group.span and nxt.merge_up:
                        vmerge = "restart"
                        break
                    n_col += nxt.span
            row_specs.append(
                GridCell(text=group.text, span=group.span, vmerge=vmerge)
            )
            col += group.span
        specs.append(row_specs)

    widths = _grid_col_widths(table, ncols)
    _set_tbl_grid(table, widths)
    tbl = table._tbl
    for tr in list(tbl.findall(qn("w:tr"))):
        tbl.remove(tr)
    for row_specs in specs:
        tr = OxmlElement("w:tr")
        col = 0
        for cell in row_specs:
            width = sum(widths[col : col + cell.span]) or 200
            tr.append(_make_grid_tc(cell.text, cell.span, width, cell.vmerge))
            col += cell.span
        tbl.append(tr)


def mark_fields_for_update(doc) -> None:
    """Ask Word to refresh TOC/page fields on open."""
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)
    else:
        existing.set(qn("w:val"), "true")
    for paragraph in doc.element.iter(qn("w:p")):
        instrs = [node.text or "" for node in paragraph.iter(qn("w:instrText"))]
        if not any("TOC" in text.upper() for text in instrs):
            continue
        for fld in paragraph.iter(qn("w:fldChar")):
            if fld.get(qn("w:fldCharType")) == "begin":
                fld.set(qn("w:dirty"), "true")


def clear_field_update_flags(doc) -> None:
    """Remove the open-time field-update prompt and TOC dirty marks."""
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is not None:
        settings.remove(existing)
    dirty_attr = qn("w:dirty")
    for paragraph in doc.element.iter(qn("w:p")):
        instrs = [node.text or "" for node in paragraph.iter(qn("w:instrText"))]
        if not any("TOC" in text.upper() for text in instrs):
            continue
        for fld in paragraph.iter(qn("w:fldChar")):
            if fld.get(qn("w:fldCharType")) == "begin" and dirty_attr in fld.attrib:
                del fld.attrib[dirty_attr]
