"""备注/Note drawer: one text box, a flat photo strip, and free-edit Excel tables."""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from PySide6.QtCore import Qt, Signal, QSize
from PySide6.QtGui import QIcon, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView,
    QButtonGroup,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QInputDialog,
    QLabel,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QStackedWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from openpyxl.utils import range_boundaries

from src.io.data_tables import (
    DataTableError,
    PreviewSnapshot,
    find_bound_row_indices,
    open_attachment,
    prepare_display_snapshot,
    read_preview_snapshot,
    sync_display_layout,
)
from src.io.project_assets import resolve_data_table_path
from src.io.test_notes import (
    add_note_images,
    create_note_workbook,
    delete_note_image,
    delete_note_table,
    list_merged_note_table_refs,
    list_note_image_paths,
    rename_note_image,
)
from src.io.test_photos import PhotoError, collect_drop_images, is_usable_test_name
from src.models.project_state import DataTableRef
from src.ui.scroll_contain import ContainedListWidget, ContainedTableWidget, ContainedTextEdit

_THUMB_SIZES = (72, 96, 128, 160)


def _local_paths(mime) -> List[Path]:
    paths = []
    for url in mime.urls():
        if url.isLocalFile():
            paths.append(Path(url.toLocalFile()))
    return paths


class _DropHost(QFrame):
    filesDropped = Signal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setObjectName("noteImageHost")

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        paths = _local_paths(event.mimeData())
        if paths:
            self.filesDropped.emit(paths)
            event.acceptProposedAction()
        else:
            event.ignore()


class _NoteGallery(ContainedListWidget):
    filesDropped = Signal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.viewport().setAcceptDrops(True)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        paths = _local_paths(event.mimeData())
        if paths:
            self.filesDropped.emit(paths)
            event.acceptProposedAction()
        else:
            event.ignore()


class TestNotePanel(QWidget):
    changed = Signal()

    def __init__(
        self,
        node_data,
        project_root: Optional[Path],
        remote_root: Optional[Path],
        leg_name: str,
        parent=None,
    ):
        super().__init__(parent)
        self.node_data = node_data
        self.project_root = Path(project_root) if project_root else None
        self.remote_root = Path(remote_root) if remote_root else None
        self.leg_name = leg_name or ""
        self._tables: List[DataTableRef] = []
        self._thumb = 96
        self._loaded = False

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(8)

        switch = QHBoxLayout()
        switch.setSpacing(8)
        self.btn_text = QPushButton("编辑文字")
        self.btn_images = QPushButton("上传图片")
        self.btn_tables = QPushButton("编辑表格")
        self._mode_buttons = (self.btn_text, self.btn_images, self.btn_tables)
        self._modes = QButtonGroup(self)
        self._modes.setExclusive(True)
        for index, button in enumerate(self._mode_buttons):
            button.setCheckable(True)
            button.setCursor(Qt.PointingHandCursor)
            self._modes.addButton(button, index)
            switch.addWidget(button)
        switch.addStretch(1)
        root.addLayout(switch)
        self._modes.idClicked.connect(self._show_mode)

        self.stack = QStackedWidget()
        self.stack.addWidget(self._build_text_page())
        self.stack.addWidget(self._build_image_page())
        self.stack.addWidget(self._build_table_page())
        root.addWidget(self.stack)

        self.btn_text.setChecked(True)
        self._apply_mode_style(0)
        self._show_mode(0)
        self._reload_text()

    def ensure_loaded(self):
        if self._loaded:
            self._reload_images()
            self._reload_tables()
            return
        self._loaded = True
        self._reload_images()
        self._reload_tables()
        self.changed.emit()

    def text(self) -> str:
        return self.editor.toPlainText()

    def image_order(self) -> List[str]:
        if not self._loaded and self.gallery.count() == 0:
            return self._saved_image_order()
        names = []
        for row in range(self.gallery.count()):
            item = self.gallery.item(row)
            name = str(item.data(Qt.UserRole) or "").strip()
            if name:
                names.append(name)
        return names

    def table_refs(self) -> List[DataTableRef]:
        if not self._loaded:
            return list(getattr(self.node_data, "note_tables", None) or [])
        return list(self._tables)

    def summary_text(self) -> str:
        bits = []
        if self.editor.toPlainText().strip():
            bits.append("有文字")
        if self._loaded:
            images = self.gallery.count()
            tables = len(self._tables)
        else:
            test_name = getattr(self.node_data, "test_name", "")
            images = len(
                list_note_image_paths(
                    self.project_root,
                    self.remote_root,
                    self.leg_name,
                    test_name,
                    order=self._saved_image_order(),
                )
            )
            tables = len(
                list_merged_note_table_refs(
                    self.project_root,
                    self.remote_root,
                    self.leg_name,
                    test_name,
                    preferred=[
                        ref.relative_path
                        for ref in (getattr(self.node_data, "note_tables", None) or [])
                    ],
                )
            )
        if images:
            bits.append(f"{images} 张")
        if tables:
            bits.append(f"{tables} 表")
        return " · ".join(bits) if bits else "未填写"

    def commit_to(self, node) -> None:
        node.note_text = self.editor.toPlainText()
        node.note_image_order = self.image_order()
        node.note_tables = list(self._tables)

    def _saved_image_order(self) -> List[str]:
        return list(getattr(self.node_data, "note_image_order", None) or [])

    def _build_text_page(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        self.editor = ContainedTextEdit()
        self.editor.setPlaceholderText("输入备注，生成报告时整段带入")
        self.editor.setMinimumHeight(140)
        self.editor.setFixedHeight(160)
        layout.addWidget(self.editor)
        return page

    def _build_image_page(self) -> QWidget:
        page = _DropHost()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(6)
        bar = QHBoxLayout()
        self.btn_pick_images = QPushButton("选择图片")
        self.btn_pick_images.clicked.connect(self._pick_images)
        self.btn_rename_image = QPushButton("改名")
        self.btn_rename_image.clicked.connect(self._rename_selected_image)
        self.btn_delete_image = QPushButton("删除")
        self.btn_delete_image.clicked.connect(self._delete_selected_image)
        self.btn_zoom_in = QPushButton("🔍+")
        self.btn_zoom_out = QPushButton("🤌-")
        self.btn_zoom_in.setObjectName("photoThumbZoom")
        self.btn_zoom_out.setObjectName("photoThumbZoom")
        self.btn_zoom_in.setFixedSize(49, 28)
        self.btn_zoom_out.setFixedSize(49, 28)
        self.btn_zoom_in.clicked.connect(lambda: self._zoom(1))
        self.btn_zoom_out.clicked.connect(lambda: self._zoom(-1))
        for button in (
            self.btn_pick_images,
            self.btn_rename_image,
            self.btn_delete_image,
            self.btn_zoom_in,
            self.btn_zoom_out,
        ):
            bar.addWidget(button)
        bar.addStretch(1)
        layout.addLayout(bar)
        self.lbl_image_hint = QLabel("可拖入 jpg / png，双击缩略图改名")
        self.lbl_image_hint.setObjectName("dimLabel")
        layout.addWidget(self.lbl_image_hint)
        self.gallery = _NoteGallery()
        self.gallery.setObjectName("noteImageList")
        self.gallery.setViewMode(QListWidget.IconMode)
        self.gallery.setFlow(QListWidget.LeftToRight)
        self.gallery.setWrapping(True)
        self.gallery.setResizeMode(QListWidget.Adjust)
        self.gallery.setMovement(QListWidget.Static)
        self.gallery.setSpacing(8)
        self.gallery.setMinimumHeight(180)
        self.gallery.itemDoubleClicked.connect(lambda _item: self._rename_selected_image())
        self.gallery.filesDropped.connect(self._import_paths)
        self._apply_thumb_size()
        layout.addWidget(self.gallery)
        page.filesDropped.connect(self._import_paths)
        return page

    def _build_table_page(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(6)
        bar = QHBoxLayout()
        self.btn_add_table = QPushButton("添加表格")
        self.btn_add_table.setToolTip("新建空白 Excel 并用本机表格程序打开")
        self.btn_add_table.clicked.connect(self._add_table)
        bar.addWidget(self.btn_add_table)
        bar.addStretch(1)
        layout.addLayout(bar)
        self.table_host = QWidget()
        self.table_layout = QVBoxLayout(self.table_host)
        self.table_layout.setContentsMargins(0, 0, 0, 0)
        self.table_layout.setSpacing(8)
        layout.addWidget(self.table_host)
        return page

    def _reload_text(self):
        self.editor.setPlainText(getattr(self.node_data, "note_text", "") or "")

    def _show_mode(self, mode: int):
        self.stack.setCurrentIndex(mode)
        self._apply_mode_style(mode)
        if mode == 1:
            self.ensure_loaded()
        elif mode == 2:
            self.ensure_loaded()

    def _apply_mode_style(self, mode: int):
        for index, button in enumerate(self._mode_buttons):
            button.setObjectName("accentButton" if index == mode else "")
            button.style().unpolish(button)
            button.style().polish(button)
            button.update()

    def _can_write(self) -> bool:
        if not is_usable_test_name(getattr(self.node_data, "test_name", "")):
            QMessageBox.warning(self, "提示", "请先选择试验名称")
            return False
        if self.project_root is None:
            QMessageBox.warning(self, "提示", "请先加载项目以确定本地镜像路径")
            return False
        if not (self.leg_name or "").strip():
            QMessageBox.warning(self, "提示", "缺少 Leg 名称")
            return False
        return True

    def _pick_images(self):
        if not self._can_write():
            return
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "选择图片",
            "",
            "图片 (*.jpg *.jpeg *.png)",
        )
        if paths:
            self._import_paths([Path(path) for path in paths])

    def _import_paths(self, paths):
        if not self._can_write():
            return
        images, skipped = collect_drop_images(paths)
        if not images:
            extra = f"\n已跳过：{', '.join(skipped[:6])}" if skipped else ""
            QMessageBox.information(self, "提示", "没有可导入的 jpg / jpeg / png。" + extra)
            return
        try:
            add_note_images(self.project_root, self.leg_name, self.node_data.test_name, images)
        except PhotoError as exc:
            QMessageBox.warning(self, "提示", str(exc))
            return
        self._reload_images()
        self.changed.emit()

    def _reload_images(self, order=None):
        if order is None:
            order = self.image_order() or self._saved_image_order()
        ordered = list_note_image_paths(
            self.project_root,
            self.remote_root,
            self.leg_name,
            getattr(self.node_data, "test_name", ""),
            order=order,
        )
        self.gallery.clear()
        for path in ordered:
            item = QListWidgetItem(path.stem)
            item.setData(Qt.UserRole, path.name)
            item.setToolTip(path.name)
            pixmap = QPixmap(str(path))
            if not pixmap.isNull():
                icon = pixmap.scaled(
                    self._thumb,
                    self._thumb,
                    Qt.KeepAspectRatio,
                    Qt.SmoothTransformation,
                )
                item.setIcon(QIcon(icon))
            self.gallery.addItem(item)
        self._apply_thumb_size()
        self._sync_zoom_buttons()

    def _selected_image_name(self) -> str:
        item = self.gallery.currentItem()
        if item is None:
            return ""
        return str(item.data(Qt.UserRole) or "")

    def _rename_selected_image(self):
        name = self._selected_image_name()
        if not name:
            QMessageBox.information(self, "提示", "请先选择一张图片")
            return
        if not self._can_write():
            return
        text, ok = QInputDialog.getText(self, "改名", "新的文件名：", text=Path(name).stem)
        if not ok:
            return
        text = (text or "").strip()
        if not text:
            return
        try:
            new = rename_note_image(
                self.project_root,
                self.remote_root,
                self.leg_name,
                self.node_data.test_name,
                name,
                text,
            )
        except PhotoError as exc:
            QMessageBox.warning(self, "提示", str(exc))
            return
        order = [new if item == name else item for item in self.image_order()]
        if new not in order:
            order.append(new)
        self._reload_images(order)
        self.changed.emit()

    def _delete_selected_image(self):
        name = self._selected_image_name()
        if not name:
            QMessageBox.information(self, "提示", "请先选择一张图片")
            return
        answer = QMessageBox.question(
            self,
            "删除图片",
            f"是否删除「{name}」？",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return
        delete_note_image(
            self.project_root,
            self.remote_root,
            self.leg_name,
            self.node_data.test_name,
            name,
        )
        order = [item for item in self.image_order() if item != name]
        self._reload_images(order)
        self.changed.emit()

    def _zoom(self, delta: int):
        try:
            index = _THUMB_SIZES.index(self._thumb)
        except ValueError:
            index = 1
        index = max(0, min(len(_THUMB_SIZES) - 1, index + delta))
        self._thumb = _THUMB_SIZES[index]
        self._reload_images(self.image_order())

    def _apply_thumb_size(self):
        self.gallery.setIconSize(QSize(self._thumb, self._thumb))
        card = self._thumb + 36
        self.gallery.setGridSize(QSize(card, card))

    def _sync_zoom_buttons(self):
        try:
            index = _THUMB_SIZES.index(self._thumb)
        except ValueError:
            index = 1
        self.btn_zoom_out.setEnabled(index > 0)
        self.btn_zoom_in.setEnabled(index < len(_THUMB_SIZES) - 1)

    def _add_table(self):
        if not self._can_write():
            return
        title, ok = QInputDialog.getText(self, "编辑表格", "表格名称：")
        if not ok:
            return
        title = (title or "").strip()
        if not title:
            QMessageBox.warning(self, "提示", "请输入表格名称")
            return
        try:
            ref = create_note_workbook(
                self.project_root, self.leg_name, self.node_data.test_name, title
            )
        except (DataTableError, PhotoError) as exc:
            QMessageBox.warning(self, "提示", str(exc))
            return
        self._reload_tables()
        self.changed.emit()
        path = resolve_data_table_path(self.project_root, self.remote_root, ref.relative_path)
        if path is not None:
            try:
                open_attachment(path)
            except DataTableError as exc:
                QMessageBox.warning(self, "无法打开", str(exc))

    def _reload_tables(self):
        preferred = [ref.relative_path for ref in self._tables] or [
            ref.relative_path for ref in (getattr(self.node_data, "note_tables", None) or [])
        ]
        self._tables = list_merged_note_table_refs(
            self.project_root,
            self.remote_root,
            self.leg_name,
            getattr(self.node_data, "test_name", ""),
            preferred=preferred,
        )
        while self.table_layout.count():
            item = self.table_layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.setParent(None)
                widget.deleteLater()
        if not self._tables:
            empty = QLabel("尚未添加表格")
            empty.setObjectName("dimLabel")
            self.table_layout.addWidget(empty)
            return
        for ref in self._tables:
            block = QFrame()
            block.setObjectName("drawerSection")
            lay = QVBoxLayout(block)
            lay.setContentsMargins(8, 8, 8, 8)
            lay.setSpacing(6)
            title = QLabel(ref.title or Path(ref.relative_path).stem)
            title.setObjectName("drawerTitle")
            lay.addWidget(title)
            actions = QHBoxLayout()
            btn_open = QPushButton("打开")
            btn_open.setToolTip("用本机 Excel / WPS / 系统默认程序打开")
            btn_refresh = QPushButton("刷新")
            btn_refresh.setToolTip("重新读取表格")
            btn_delete = QPushButton("删除")
            btn_open.clicked.connect(lambda _=False, r=ref: self._open_table(r))
            btn_refresh.clicked.connect(lambda _=False, r=ref: self._refresh_table(r))
            btn_delete.clicked.connect(lambda _=False, r=ref: self._delete_table(r))
            actions.addWidget(btn_open)
            actions.addWidget(btn_refresh)
            actions.addWidget(btn_delete)
            actions.addStretch(1)
            lay.addLayout(actions)
            preview = ContainedTableWidget(0, 0)
            preview.setEditTriggers(QAbstractItemView.NoEditTriggers)
            preview.setMinimumHeight(120)
            preview.setMaximumHeight(220)
            preview.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            snap = self._read_snapshot(ref)
            if snap is not None:
                self._fill_preview(preview, snap)
            lay.addWidget(preview)
            self.table_layout.addWidget(block)

    def _read_snapshot(self, ref: DataTableRef) -> Optional[PreviewSnapshot]:
        path = resolve_data_table_path(self.project_root, self.remote_root, ref.relative_path)
        if path is None:
            return None
        try:
            return read_preview_snapshot(path)
        except Exception:
            return None

    def _fill_preview(self, table: ContainedTableWidget, snap: PreviewSnapshot):
        display = prepare_display_snapshot(snap, include_limit_row=True)
        table.clearSpans()
        table.clear()
        rows = display.values or []
        cols = max((len(row) for row in rows), default=0)
        table.setRowCount(len(rows))
        table.setColumnCount(cols)
        bound_rows = set(find_bound_row_indices(display).indices())
        for r, row in enumerate(rows):
            for c in range(cols):
                text = row[c] if c < len(row) else ""
                cell = QTableWidgetItem(text)
                cell.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                if c == 0 or (c > 0 and r in bound_rows):
                    cell.setTextAlignment(Qt.AlignCenter)
                table.setItem(r, c, cell)
        origin_r = display.origin_row or 1
        origin_c = display.origin_col or 1
        for merge in display.merges or []:
            try:
                min_c, min_r, max_c, max_r = range_boundaries(merge)
            except Exception:
                continue
            r0 = min_r - origin_r
            c0 = min_c - origin_c
            r1 = max_r - origin_r
            c1 = max_c - origin_c
            if r0 < 0 or c0 < 0 or r1 >= len(rows) or c1 >= cols:
                continue
            row_span = r1 - r0 + 1
            col_span = c1 - c0 + 1
            if row_span > 1 or col_span > 1:
                table.setSpan(r0, c0, row_span, col_span)

    def _open_table(self, ref: DataTableRef):
        path = resolve_data_table_path(self.project_root, self.remote_root, ref.relative_path)
        if path is None:
            QMessageBox.warning(self, "提示", "找不到备注表格（本地与公盘均无）")
            return
        try:
            open_attachment(path)
        except DataTableError as exc:
            QMessageBox.warning(self, "无法打开", str(exc))

    def _refresh_table(self, ref: DataTableRef):
        path = resolve_data_table_path(self.project_root, self.remote_root, ref.relative_path)
        if path is None:
            QMessageBox.warning(self, "提示", "找不到备注表格（本地与公盘均无）")
            return
        try:
            sync_display_layout(path)
        except DataTableError as exc:
            QMessageBox.warning(self, "提示", str(exc))
            return
        self._reload_tables()
        self.changed.emit()

    def _delete_table(self, ref: DataTableRef):
        answer = QMessageBox.question(
            self,
            "删除表格",
            f"是否删除「{ref.title}」？",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return
        try:
            delete_note_table(self.project_root, self.remote_root, ref.relative_path)
        except DataTableError as exc:
            QMessageBox.warning(self, "提示", str(exc))
            return
        self._tables = [item for item in self._tables if item.relative_path != ref.relative_path]
        self._reload_tables()
        self.changed.emit()
